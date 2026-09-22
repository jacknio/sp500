"""Treasury yield / stock-bond regime study -- ONE standalone file to copy.

Requirements: Python 3.10+ and a Bloomberg Terminal with Desktop API access.
Dependencies: python -m pip install numpy pandas matplotlib openpyxl
blpapi: use your firm's Bloomberg Python environment / official SDK installer.

First run and subsequent updates:
    python treasury_regime_study.py fetch
    python treasury_regime_study.py analyze
Combined command: python treasury_regime_study.py run
Calculation check: python treasury_regime_study.py check
Historical cutoff: python treasury_regime_study.py run --asof 2026-09-14
Optional images: python treasury_regime_study.py analyze --images

FLOW: Bloomberg -> raw/<timestamp>/observations.csv + manifest.json
      -> read from disk -> calculate -> output/<timestamp>/research.xlsx
      + run_manifest.json; --images additionally exports charts/*.svg/png/pdf
No HTML, template, other .py files or existing dashboard is needed.
Review map: section 1 = config; 2 = fetch/read; 3-4 = math; 5 = outputs.

Transparent reconstruction: the author's undisclosed sample/parameters may differ.
Charts and findings are calculated from actual inputs, not hardcoded conclusions.
check uses isolated test values only; it never writes to raw/ or output/.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import sys
import tempfile
import time
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import numpy as np
import pandas as pd


# ============================================================================
# SECTION 1. EDIT HERE: Bloomberg mappings and research assumptions.
# Raw values are NEVER rescaled on download. unit and multiplier are explicit.
# LF98OAS is quoted in percentage points: 4.34 = 434 bp, NOT 4.34 bp.
# ============================================================================
ROOT = Path(__file__).resolve().parent
VERSION = "1.0"
HOST, PORT, TIMEOUT = "localhost", 8194, 120
START = "1973-01-01"
THRESHOLD = 5.25
FORWARD_WEEKS = 13
CORRELATION_WEEKS = 104
CORRELATION_COLUMN = f"corr_{CORRELATION_WEEKS}w"
TOTAL_RETURN_COLUMN = f"corr_total_return_{CORRELATION_WEEKS}w"
MIN_TREND_MONTHS = 120
BOOTSTRAPS = 500

# key: (security, field, request frequency, raw unit, required)
SERIES = {
    "us10y": ("USGG10YR Index", "PX_LAST", "DAILY", "percent", True),
    "ust": ("LUATTRUU Index", "PX_LAST", "DAILY", "total return index", True),
    "ust_monthly": ("LUATTRUU Index", "PX_LAST", "MONTHLY", "total return index", True),
    "spx": ("SPX Index", "PX_LAST", "DAILY", "price index", True),
    "sptr": ("SPTR Index", "PX_LAST", "DAILY", "total return index", False),
    "move": ("MOVE Index", "PX_LAST", "DAILY", "index", False),
    "liquidity": ("GVLQUSD Index", "PX_LAST", "DAILY", "index", False),
    "vix": ("VIX Index", "PX_LAST", "DAILY", "index", False),
    "hy_oas": ("LF98OAS Index", "PX_LAST", "DAILY", "percent", False),
    "bcom": ("BCOM Index", "PX_LAST", "DAILY", "index", False),
}
RAW_COLUMNS = ["series", "security", "field", "frequency", "date", "value", "unit", "retrieved_at"]
CORE = ["us10y", "ust", "spx"]
SOURCES = {
    "Treasury index mapping": "https://assets.bbhub.io/professional/sites/10/999212_Fixed-income-investment-insight-hedging-inflation-toolkit.pdf",
    "Bloomberg API": "https://bloomberg.github.io/blpapi-docs/python/3.26.9/_autosummary/blpapi.Session.html",
    "HY OAS quote units (4.34 percent)": "https://doubleline.com/wp-content/uploads/12-6-2022-TR-Webcast-FINAL.pdf",
    "Liquidity index mapping (research paper)": "https://www.sciencedirect.com/science/article/abs/pii/S1572308924001542",
}


def atomic_text(path, text):
    """Close before replace for Windows; keep the old file on failed writes."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile("w", encoding="utf-8", newline="", dir=path.parent,
                                         delete=False, suffix=".tmp") as stream:
            temporary = Path(stream.name)
            stream.write(text)
        temporary.replace(path)
    finally:
        if temporary:
            temporary.unlink(missing_ok=True)


def save_json(path, obj):
    atomic_text(path, json.dumps(obj, indent=2, ensure_ascii=False, allow_nan=False))


def timestamp():
    return datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")


# ============================================================================
# SECTION 2. FETCH / READ. Only fully received series enter committed snapshots.
# No synthetic data, previous-value fill, silent alternate ticker, or API secrets.
# ============================================================================
def parse_response(message, key, retrieved, start, end):
    """One security per request; accept both partial and final response messages."""
    if message.hasElement("responseError"):
        raise ValueError(str(message.getElement("responseError")))
    data = message.getElement("securityData")
    if data.hasElement("securityError"):
        raise ValueError(str(data.getElement("securityError")))
    if data.hasElement("fieldExceptions") and data.getElement("fieldExceptions").numValues():
        raise ValueError(str(data.getElement("fieldExceptions")))
    security, field, frequency, unit, _ = SERIES[key]
    if " ".join(data.getElementAsString("security").split()).casefold() != security.casefold():
        raise ValueError("Unexpected security in response for " + key)
    values, rows = data.getElement("fieldData"), []
    for i in range(values.numValues()):
        row = values.getValueAsElement(i)
        day = row.getElementAsDatetime("date")
        day = day.date() if isinstance(day, datetime) else day
        if not isinstance(day, date) or not start <= day <= end:
            raise ValueError("Bloomberg date outside request")
        value = None
        if row.hasElement(field) and not row.getElement(field).isNull():
            value = row.getElementAsFloat(field)
            if not math.isfinite(value):
                raise ValueError("Nonfinite Bloomberg value")
        rows.append([key, security, field, frequency, day.isoformat(), value, unit, retrieved])
    return rows


def fetch(start, end, root=ROOT):
    """Request daily observations AND monthly Treasury history for the YoY chart."""
    try:
        import blpapi
    except ImportError as exc:
        raise RuntimeError("blpapi unavailable. Run fetch in your company's Bloomberg-enabled Python environment.") from exc
    folder = root / "raw" / timestamp()
    folder.mkdir(parents=True)
    retrieved = datetime.now(timezone.utc).isoformat()
    manifest = dict(version=VERSION, status="running", start=str(start), end=str(end),
                    retrieved_at=retrieved, series=SERIES, errors={}, requests={})
    rows, session = [], None
    save_json(folder / "manifest.json", manifest)
    try:
        options = blpapi.SessionOptions()
        options.setServerHost(HOST)
        options.setServerPort(PORT)
        options.setConnectTimeout(10000)
        session = blpapi.Session(options)
        if not session.start() or not session.openService("//blp/refdata"):
            raise ConnectionError("Cannot connect to Desktop API localhost:8194. Check Terminal login and API permissions.")
        service = session.getService("//blp/refdata")
        for number, (key, cfg) in enumerate(SERIES.items(), 1):
            security, field, frequency, _, required = cfg
            print(f"FETCH {key}: {security} / {field} / {frequency}", flush=True)
            correlation = blpapi.CorrelationId(number)
            finished, sent = False, False
            try:
                request = service.createRequest("HistoricalDataRequest")
                request.getElement("securities").appendValue(security)
                request.getElement("fields").appendValue(field)
                settings = dict(startDate=start.strftime("%Y%m%d"), endDate=end.strftime("%Y%m%d"),
                                periodicitySelection=frequency, periodicityAdjustment="ACTUAL",
                                nonTradingDayFillOption="ACTIVE_DAYS_ONLY", nonTradingDayFillMethod="NIL_VALUE")
                for setting, value in settings.items():
                    request.set(setting, value)
                manifest["requests"][key] = settings
                session.sendRequest(request, correlationId=correlation)
                sent, own_rows = True, []
                deadline = time.monotonic() + TIMEOUT
                while not finished:
                    if time.monotonic() >= deadline:
                        raise TimeoutError(f"{key} timed out after {TIMEOUT}s")
                    event = session.nextEvent(1000)
                    for message in event:
                        if str(message.messageType()) in ("SessionTerminated", "SessionStartupFailure", "SessionConnectionDown"):
                            raise ConnectionError(str(message))
                        if not any(c.value() == number for c in message.correlationIds()):
                            continue
                        with (folder / f"{key}_response.txt").open("a", encoding="utf-8") as log:
                            log.write(str(message) + "\n")
                        if event.eventType() == blpapi.Event.REQUEST_STATUS:
                            if str(message.messageType()) == "RequestFailure":
                                raise RuntimeError(str(message))
                            continue
                        if event.eventType() not in (blpapi.Event.PARTIAL_RESPONSE, blpapi.Event.RESPONSE):
                            continue
                        own_rows.extend(parse_response(message, key, retrieved, start, end))
                        finished = event.eventType() == blpapi.Event.RESPONSE
                if not any(row[5] is not None for row in own_rows):
                    raise ValueError("No numeric observations")
                if len({row[4] for row in own_rows}) != len(own_rows):
                    raise ValueError("Duplicate observation dates")
                rows.extend(own_rows)
            except Exception as exc:
                manifest["errors"][key] = str(exc)
                if sent and not finished:
                    try:
                        session.cancel(correlation)
                    except Exception as cancel_error:
                        manifest["errors"][key] += f"; cancel error: {cancel_error}"
                if required or isinstance(exc, ConnectionError):
                    raise
                print(f"  Optional series unavailable: {exc}", flush=True)
            # Checkpoint only fully validated responses. Failed partial data stay in logs.
            atomic_text(folder / "observations.csv", pd.DataFrame(rows, columns=RAW_COLUMNS).to_csv(index=False))
            save_json(folder / "manifest.json", manifest)
        validate_raw(pd.DataFrame(rows, columns=RAW_COLUMNS), manifest)
        manifest["status"] = "complete"
        manifest["sha256"] = hashlib.sha256((folder / "observations.csv").read_bytes()).hexdigest()
        save_json(folder / "manifest.json", manifest)
        # Single commit point: failed runs never replace the latest successful snapshot.
        save_json(root / "raw" / "latest.json", {"snapshot": folder.name})
    except BaseException as exc:
        manifest.update(status="failed", run_error=str(exc))
        atomic_text(folder / "observations.csv", pd.DataFrame(rows, columns=RAW_COLUMNS).to_csv(index=False))
        save_json(folder / "manifest.json", manifest)
        raise
    finally:
        if session is not None:
            try:
                session.stop()
            except Exception as stop_error:
                print(f"WARNING: Bloomberg session cleanup failed: {stop_error}", file=sys.stderr)
    print("Raw data saved:", folder)
    return folder


def validate_raw(raw, manifest):
    if list(raw.columns) != RAW_COLUMNS or raw.empty:
        raise ValueError("Raw CSV is empty or its schema changed")
    if raw[[c for c in RAW_COLUMNS if c != "value"]].isna().any().any():
        raise ValueError("Missing raw metadata")
    raw = raw.copy()
    raw["date"] = pd.to_datetime(raw["date"], format="%Y-%m-%d", errors="raise")
    raw["value"] = pd.to_numeric(raw["value"], errors="raise")
    if np.isinf(raw["value"]).any() or raw.duplicated(["series", "date"]).any():
        raise ValueError("Infinite values or duplicate series/date pairs")
    if not raw["date"].between(pd.Timestamp(manifest["start"]), pd.Timestamp(manifest["end"])).all():
        raise ValueError("CSV dates outside snapshot range")
    for key, group in raw.groupby("series"):
        if key not in SERIES or list(SERIES[key]) != list(manifest["series"].get(key, [])):
            raise ValueError(f"{key}: configuration differs from saved snapshot. Restore mappings or fetch again.")
        for column, expected in zip(["security", "field", "frequency", "unit"], SERIES[key][:4]):
            if not group[column].eq(expected).all():
                raise ValueError(f"{key}: unexpected {column}")
        if not group["retrieved_at"].eq(manifest["retrieved_at"]).all():
            raise ValueError("Mixed raw snapshots")
        values = group["value"].dropna()
        if key in ("ust", "ust_monthly", "spx", "sptr", "bcom") and values.le(0).any():
            raise ValueError(f"{key}: index levels must be positive")
        if key == "us10y" and not values.between(-5, 40).all():
            raise ValueError("10Y yield must be in percent (5.25), not basis points (525)")
    for key, cfg in SERIES.items():
        if cfg[-1] and raw.loc[raw["series"].eq(key), "value"].dropna().empty:
            raise ValueError(f"Required history missing: {key}")
    return raw.sort_values(["series", "date"]).reset_index(drop=True)


def read_snapshot(root=ROOT, snapshot=None):
    if snapshot is None:
        pointer = root / "raw" / "latest.json"
        if not pointer.exists():
            raise FileNotFoundError("No Bloomberg snapshot. Run: python treasury_regime_study.py fetch")
        name = json.loads(pointer.read_text(encoding="utf-8"))["snapshot"]
        if Path(name).name != name:
            raise ValueError("Invalid snapshot pointer")
        snapshot = root / "raw" / name
    snapshot = Path(snapshot).resolve()
    manifest = json.loads((snapshot / "manifest.json").read_text(encoding="utf-8"))
    if manifest["status"] != "complete":
        raise ValueError("Snapshot is incomplete; see manifest errors")
    path = snapshot / "observations.csv"
    if hashlib.sha256(path.read_bytes()).hexdigest() != manifest["sha256"]:
        raise ValueError("Raw CSV hash mismatch. Do not manually overwrite downloaded observations.")
    return validate_raw(pd.read_csv(path, float_precision="round_trip"), manifest), manifest


# ============================================================================
# SECTION 3. TIME ALIGNMENT / MATH. Fractions internally; percentages in tables.
# Missing weeks remain missing. No fill and no interpolated monthly history.
# ============================================================================
def weekly_sample(daily, end):
    """Sample a COMPLETE same-date row, not each column's independent last value.

    W-FRI is the observation bucket; source_date is the actual common close.
    Discard partial final weeks. Holiday Thursday closes stay dated Thursday.
    """
    grid = pd.date_range(daily.index.min(), pd.Timestamp(end), freq="W-FRI")
    valid = daily.dropna(how="any").copy()
    valid["source_date"] = valid.index
    if valid.empty:
        return pd.DataFrame(index=grid, columns=[*daily.columns, "source_date"])
    sampled = valid.resample("W-FRI").last().reindex(grid)
    # A week can never borrow a value from an earlier week.
    age = (sampled.index.to_series() - sampled["source_date"]).dt.days
    sampled.loc[age.gt(4)] = np.nan
    return sampled


def valid_correlation(x, y, window):
    result = x.rolling(window, min_periods=window).corr(y)
    ok = x.rolling(window).std().gt(1e-14) & y.rolling(window).std().gt(1e-14)
    return result.where(ok).clip(-1, 1)


def forward_compound(returns, horizon):
    """At t: product of (1+r[t+1] ... 1+r[t+h])-1; all h weeks required."""
    return np.expm1(np.log1p(returns).rolling(horizon, min_periods=horizon).sum().shift(-horizon))


def trend_analysis(monthly):
    """Calendar 12M TR growth; OLS on month number, residual SD sqrt(SSE/(n-2)).

    Full-sample fit is only descriptive. prior_* uses observations strictly < t.
    ±SD lines are residual reference bands, not confidence / forecast intervals.
    """
    result = pd.DataFrame({"ust_index": monthly})
    result["yoy_pct"] = monthly.pct_change(12, fill_method=None) * 100
    x = np.arange(len(result), dtype=float) / 12
    y = result["yoy_pct"].to_numpy(dtype=float)
    result["trend_pct"] = np.nan
    for name in ("minus_2sd", "minus_1sd", "plus_1sd", "plus_2sd", "prior_trend_pct", "prior_z"):
        result[name] = np.nan
    mask = np.isfinite(y)
    if mask.sum() >= 24:
        design = np.column_stack([np.ones(mask.sum()), x[mask]])
        intercept, slope = np.linalg.lstsq(design, y[mask], rcond=None)[0]
        fit = intercept + slope * x
        sd = np.sqrt(np.sum((y[mask] - fit[mask]) ** 2) / (mask.sum() - 2))
        result["trend_pct"] = fit
        for name, k in (("minus_2sd", -2), ("minus_1sd", -1), ("plus_1sd", 1), ("plus_2sd", 2)):
            result[name] = fit + k * sd
    for t in range(len(result)):
        prior = np.isfinite(y[:t])
        if prior.sum() < MIN_TREND_MONTHS or not np.isfinite(y[t]):
            continue
        design = np.column_stack([np.ones(prior.sum()), x[:t][prior]])
        coefficients = np.linalg.lstsq(design, y[:t][prior], rcond=None)[0]
        fitted = float(np.array([1, x[t]]) @ coefficients)
        sd = np.sqrt(np.sum((y[:t][prior] - design @ coefficients) ** 2) / (prior.sum() - 2))
        result.iloc[t, result.columns.get_loc("prior_trend_pct")] = fitted
        if sd > 1e-10:
            result.iloc[t, result.columns.get_loc("prior_z")] = (y[t] - fitted) / sd
    return result


def calculate(raw, end):
    """Keep core market dates separate from optional index coverage."""
    raw = raw.loc[raw["date"].le(pd.Timestamp(end))].copy()
    if raw.empty:
        raise ValueError("No data on or before --asof")
    daily = raw.loc[raw["frequency"].eq("DAILY")].pivot(index="date", columns="series", values="value")
    weekly = weekly_sample(daily[CORE], end)
    weekly["stock_return"] = weekly["spx"].pct_change(fill_method=None)
    weekly["bond_return"] = weekly["ust"].pct_change(fill_method=None)
    for window in sorted({52, CORRELATION_WEEKS, 156}):
        weekly[f"corr_{window}w"] = valid_correlation(weekly["stock_return"], weekly["bond_return"], window)
    weekly["yield_change_13w_bp"] = (weekly["us10y"] - weekly["us10y"].shift(13)) * 100
    weekly["yield_vol_26w_bp"] = (weekly["us10y"].diff() * 100).rolling(26).std() * np.sqrt(52)
    weekly["portfolio_return"] = 0.6 * weekly["stock_return"] + 0.4 * weekly["bond_return"]
    for key in ("stock", "bond", "portfolio"):
        weekly[f"forward_{key}_pct"] = forward_compound(weekly[f"{key}_return"], FORWARD_WEEKS) * 100
    # A positive correlation is not the same event as losing on both assets.
    complete = weekly[["forward_stock_pct", "forward_bond_pct"]].notna().all(axis=1)
    weekly["forward_joint_loss"] = (weekly["forward_stock_pct"].lt(0) & weekly["forward_bond_pct"].lt(0)).astype(float).where(complete)
    for key in ("move", "liquidity", "vix", "hy_oas", "bcom", "sptr"):
        if key not in daily:
            continue
        sampled = weekly_sample(daily[[key]], end).reindex(weekly.index)
        weekly[key] = sampled[key] * (100 if key == "hy_oas" else 1)
        weekly[f"{key}_date"] = sampled["source_date"]
        weekly[f"{key}_13w_average"] = weekly[key].rolling(13).mean()
        # Only same-date observations are used for conditional forward changes.
        aligned = pd.Series(daily[key].reindex(pd.DatetimeIndex(weekly["source_date"])).to_numpy(), index=weekly.index)
        aligned *= 100 if key == "hy_oas" else 1
        weekly[f"{key}_at_core_close"] = aligned
        weekly[f"forward_{key}_change"] = aligned.shift(-FORWARD_WEEKS) - aligned
    if "sptr" in daily:
        weekly[TOTAL_RETURN_COLUMN] = valid_correlation(weekly["sptr_at_core_close"].pct_change(fill_method=None),
                                                        weekly["bond_return"], CORRELATION_WEEKS)

    # Sensitivity to the unspecified sampling convention: 5 matched sessions,
    # 504 observations. These overlap; they are NOT 504 independent weeks.
    common = daily[CORE].dropna().copy()
    changes = common[["spx", "ust"]].pct_change(5, fill_method=None)
    gap = common.index.to_series().diff(5).dt.days
    changes.loc[~gap.between(5, 10)] = np.nan
    common["corr_504d_5session"] = valid_correlation(changes["spx"], changes["ust"], 504)
    common["window_calendar_days"] = common.index.to_series().diff(503).dt.days
    daily_sensitivity = common[["us10y", "corr_504d_5session", "window_calendar_days"]]

    m = raw.loc[raw["series"].eq("ust_monthly")].set_index("date")["value"].dropna()
    m = m.groupby(m.index.to_period("M")).last()
    if m.empty:
        raise ValueError("No monthly Treasury history at cutoff")
    month_grid = pd.period_range(m.index.min(), pd.Timestamp(end).to_period("M"), freq="M")
    monthly = m.reindex(month_grid)
    monthly.index = monthly.index.to_timestamp(how="end").normalize()
    monthly = monthly.loc[monthly.index <= pd.Timestamp(end)]  # Exclude partial month.
    if monthly.pct_change(12, fill_method=None).dropna().empty:
        raise ValueError("At least two matching month-end values 12 months apart are required for the Treasury YoY chart")
    trend = trend_analysis(monthly)
    # Core series coverage by year makes monthly-only early histories visible.
    coverage = raw.assign(year=raw["date"].dt.year).groupby(["series", "year"])["value"].count().rename("source_observations").reset_index()
    if weekly[CORRELATION_COLUMN].dropna().empty:
        raise ValueError(f"No full {CORRELATION_WEEKS}-week correlation window. Check daily source history, missing weeks and requested date range.")
    return weekly, trend, daily_sensitivity, coverage


# ============================================================================
# SECTION 4. RESEARCH TESTS. Replication, robustness and forward outcomes.
# ============================================================================
def bucket_table(frame, corr=CORRELATION_COLUMN):
    """[lower, upper); zero correlation remains in the denominator and is shown."""
    edges = [-np.inf, *np.arange(3.5, 7.5001, 0.25), np.inf]
    rows = []
    valid = frame[["us10y", corr]].notna().all(axis=1)
    for lower, upper in zip(edges[:-1], edges[1:]):
        selected = valid & frame["us10y"].ge(lower) & frame["us10y"].lt(upper)
        values = frame.loc[selected, corr]
        label = f"{lower:g}–{upper:g}%" if np.isfinite([lower, upper]).all() else ("<3.5%" if lower < 0 else ">=7.5%")
        n = len(values)
        rows.append(dict(bucket=label, n=n, positive_pct=100 * values.gt(0).mean() if n else np.nan,
                         negative_pct=100 * values.lt(0).mean() if n else np.nan,
                         zero_pct=100 * values.eq(0).mean() if n else np.nan,
                         median_correlation=values.median(),
                         first=str(values.index.min().date()) if n else "",
                         last=str(values.index.max().date()) if n else "",
                         calendar_years=values.index.year.nunique(),
                         contiguous_runs=int((selected & ~selected.shift(fill_value=False)).sum()),
                         pre2000_pct=100 * (values.index.year < 2000).mean() if n else np.nan))
    return pd.DataFrame(rows)


def conditional_summary(frame, threshold=THRESHOLD, corr=CORRELATION_COLUMN):
    rows = []
    for period, eligible in (("All", np.ones(len(frame), dtype=bool)),
                             ("Before 2000", frame.index.year < 2000),
                             ("2000–2019", (frame.index.year >= 2000) & (frame.index.year < 2020)),
                             ("2020 onward", frame.index.year >= 2020)):
        for high in (False, True):
            mask = pd.Series(eligible, index=frame.index) & frame["us10y"].notna()
            mask &= frame["us10y"].ge(threshold) if high else frame["us10y"].lt(threshold)
            observed = frame.loc[mask, corr].dropna()
            futures = frame.loc[mask, ["forward_stock_pct", "forward_bond_pct", "forward_joint_loss"]].dropna()
            rows.append(dict(period=period, regime=f">= {threshold:g}%" if high else f"< {threshold:g}%",
                             correlation_n=len(observed), positive_correlation_pct=100 * observed.gt(0).mean(),
                             forward_n=len(futures), joint_loss_pct=100 * futures["forward_joint_loss"].mean(),
                             stock_median_pct=futures["forward_stock_pct"].median(), bond_median_pct=futures["forward_bond_pct"].median()))
    return pd.DataFrame(rows)


def block_bootstrap_difference(yields, outcome, threshold, block, repetitions):
    """Paired circular moving-block bootstrap; preserve regular calendar gaps.

    CI is exploratory and conditional on this historical sample. It does not
    correct threshold selection or establish stability across structural breaks.
    """
    y, o = np.asarray(yields, float), np.asarray(outcome, float)
    valid = np.isfinite(y) & np.isfinite(o)
    def estimate(index):
        hi = valid[index] & (y[index] >= threshold)
        lo = valid[index] & (y[index] < threshold)
        if hi.sum() < 20 or lo.sum() < 20:
            return np.nan
        return 100 * (o[index][hi].mean() - o[index][lo].mean())
    observed = estimate(np.arange(len(y)))
    if not np.isfinite(observed) or len(y) < 2 * block:
        return observed, np.nan, np.nan, 0
    rng = np.random.default_rng(20260915)
    samples = []
    for _ in range(repetitions):
        starts = rng.integers(0, len(y), size=math.ceil(len(y) / block))
        indices = ((starts[:, None] + np.arange(block)) % len(y)).ravel()[:len(y)]
        estimate_value = estimate(indices)
        if np.isfinite(estimate_value):
            samples.append(estimate_value)
    if len(samples) < max(100, 0.8 * repetitions):
        return observed, np.nan, np.nan, len(samples)
    low, high = np.percentile(samples, [2.5, 97.5])
    return observed, low, high, len(samples)


def deeper_research(weekly, daily_sensitivity, repetitions):
    """Pre-specified sensitivity tables; no search for an 'optimal' yield level."""
    rows = []
    for threshold in np.arange(4.0, 6.5001, 0.25):
        valid = weekly[CORRELATION_COLUMN].notna() & weekly["us10y"].notna()
        high, low = valid & weekly["us10y"].ge(threshold), valid & weekly["us10y"].lt(threshold)
        rows.append(dict(threshold_pct=threshold, high_n=int(high.sum()), low_n=int(low.sum()),
                         high_positive_pct=100 * weekly.loc[high, CORRELATION_COLUMN].gt(0).mean(),
                         low_positive_pct=100 * weekly.loc[low, CORRELATION_COLUMN].gt(0).mean()))
    thresholds = pd.DataFrame(rows)
    bootstrap = []
    for outcome in ("positive_correlation", "forward_joint_loss"):
        values = (weekly[CORRELATION_COLUMN].gt(0).astype(float).where(weekly[CORRELATION_COLUMN].notna())
                  if outcome == "positive_correlation" else weekly[outcome])
        for block in (26, 52, 104):
            effect, low, high, n = block_bootstrap_difference(weekly["us10y"], values, THRESHOLD, block, repetitions)
            bootstrap.append(dict(outcome=outcome, block_weeks=block, difference_pp=effect,
                                  ci95_low_pp=low, ci95_high_pp=high, valid_bootstraps=n))

    # Require four preceding weeks below threshold; then prevent overlapping
    # Forward event windows. The configured horizon must have fully matured.
    below = weekly["us10y"].lt(THRESHOLD) & weekly["us10y"].notna()
    crossing = weekly["us10y"].ge(THRESHOLD) & below.shift(1).rolling(4).sum().eq(4)
    events, last_position = [], -FORWARD_WEEKS - 1
    forward_cols = [c for c in weekly if c.startswith("forward_")]
    for position in np.flatnonzero(crossing):
        if position - last_position <= FORWARD_WEEKS:
            continue
        row = weekly.iloc[position]
        if pd.isna(row["forward_joint_loss"]):
            continue
        last_position = position
        events.append(dict(date=weekly.index[position], source_date=row["source_date"],
                           yield_pct=row["us10y"], **{c: row[c] for c in forward_cols}))
    event_table = pd.DataFrame(events, columns=["date", "source_date", "yield_pct", *forward_cols])

    # All phase offsets, each with non-overlapping forward-horizon windows.
    # Do not cherry-pick the phase that most supports the article.
    phases = []
    for phase in range(FORWARD_WEEKS):
        f = weekly.iloc[phase::FORWARD_WEEKS]
        for high in (False, True):
            selected = f["us10y"].ge(THRESHOLD) if high else f["us10y"].lt(THRESHOLD)
            outcomes = f.loc[selected, "forward_joint_loss"].dropna()
            phases.append(dict(phase=phase, regime="High yield" if high else "Lower yield", n=len(outcomes),
                               joint_loss_pct=100 * outcomes.mean()))
    methods = []
    method_columns = list(dict.fromkeys(["corr_52w", CORRELATION_COLUMN, "corr_156w", TOTAL_RETURN_COLUMN]))
    for column in method_columns:
        if column not in weekly:
            continue
        for high in (False, True):
            selection = weekly["us10y"].ge(THRESHOLD) if high else weekly["us10y"].lt(THRESHOLD)
            values = weekly.loc[selection, column].dropna()
            methods.append(dict(sample="Own history", method=column, regime="High yield" if high else "Lower yield", n=len(values), positive_pct=100 * values.gt(0).mean(),
                                first=str(values.index.min().date()) if len(values) else "", last=str(values.index.max().date()) if len(values) else ""))
    for high in (False, True):
        selection = daily_sensitivity["us10y"].ge(THRESHOLD) if high else daily_sensitivity["us10y"].lt(THRESHOLD)
        values = daily_sensitivity.loc[selection, "corr_504d_5session"].dropna()
        methods.append(dict(sample="Own history (daily)", method="corr_504d_5session", regime="High yield" if high else "Lower yield", n=len(values), positive_pct=100 * values.gt(0).mean(),
                            first=str(values.index.min().date()) if len(values) else "", last=str(values.index.max().date()) if len(values) else "",
                            median_window_days=daily_sensitivity.loc[values.index,"window_calendar_days"].median()))

    # Same weekly dates separate sampling effects from differences in coverage.
    comparable = weekly[[c for c in method_columns if c in weekly and weekly[c].notna().any()]].copy()
    aligned_daily = daily_sensitivity["corr_504d_5session"].reindex(pd.DatetimeIndex(weekly["source_date"])).to_numpy()
    if np.isfinite(aligned_daily).any():
        comparable["corr_504d_5session"] = aligned_daily
    comparable = comparable.dropna()
    for column in comparable:
        for high in (False, True):
            yields = weekly.loc[comparable.index,"us10y"]
            values = comparable.loc[yields.ge(THRESHOLD) if high else yields.lt(THRESHOLD),column]
            methods.append(dict(sample="Common weekly dates", method=column, regime="High yield" if high else "Lower yield", n=len(values), positive_pct=100 * values.gt(0).mean(),
                                first=str(values.index.min().date()) if len(values) else "", last=str(values.index.max().date()) if len(values) else ""))

    # Is it the level or the recent direction of rates? Observable states at t.
    regimes = []
    for high in (False, True):
        for rising in (False, True):
            mask = weekly["us10y"].ge(THRESHOLD) if high else weekly["us10y"].lt(THRESHOLD)
            mask &= weekly["yield_change_13w_bp"].gt(0) if rising else weekly["yield_change_13w_bp"].le(0)
            out = weekly.loc[mask, ["forward_stock_pct", "forward_bond_pct", "forward_joint_loss"]].dropna()
            regimes.append(dict(yield_level="High" if high else "Lower", direction="Rising" if rising else "Flat / falling",
                                n=len(out), joint_loss_pct=100 * out["forward_joint_loss"].mean(),
                                stock_median_pct=out["forward_stock_pct"].median(), bond_median_pct=out["forward_bond_pct"].median()))
    return dict(thresholds=thresholds, bootstrap=pd.DataFrame(bootstrap), events=event_table,
                nonoverlap_phases=pd.DataFrame(phases), sampling_methods=pd.DataFrame(methods), yield_regimes=pd.DataFrame(regimes))


def transmission_tables(raw, end):
    """Use the same actual date within each pair; compare levels AND changes."""
    daily = raw.loc[raw["date"].le(pd.Timestamp(end)) & raw["frequency"].eq("DAILY")].pivot(index="date", columns="series", values="value")
    results, pairs = [], {}
    for first, second in (("move", "liquidity"), ("move", "vix"), ("vix", "hy_oas")):
        if first not in daily or second not in daily:
            continue
        pair = weekly_sample(daily[[first, second]], end)[[first, second]]
        if "hy_oas" in pair:
            pair["hy_oas"] *= 100
        pairs[f"{first}_{second}"] = pair
        for period, mask in (("All", np.ones(len(pair), dtype=bool)),
                             ("Before 2020", pair.index.year < 2020),
                             ("2020 onward", pair.index.year >= 2020)):
            for measure, values in (("Levels", pair), ("1W changes", pair.diff())):
                data = values.loc[mask].dropna()
                corr = data[first].corr(data[second]) if len(data) >= 20 and data.std().gt(1e-12).all() else np.nan
                results.append(dict(pair=f"{first} / {second}", period=period, measure=measure,
                                    n=len(data), correlation=corr))
    return pd.DataFrame(results, columns=["pair", "period", "measure", "n", "correlation"]), pairs


def forward_channels(weekly):
    """Directly test the article's subsequent volatility/spread claim, descriptively.

    Each outcome has its own sample count. No missing change becomes a zero.
    These overlapping forward changes do not establish a causal transmission chain.
    """
    rows = []
    for key, unit in (("move", "MOVE points"), ("vix", "VIX points"),
                      ("hy_oas", "basis points"), ("liquidity", "liquidity index points")):
        column = f"forward_{key}_change"
        if column not in weekly:
            continue
        for period, era in (("All", weekly.index.year > 0), ("Before 2000", weekly.index.year < 2000),
                            ("2000–2019", (weekly.index.year >= 2000) & (weekly.index.year < 2020)),
                            ("2020 onward", weekly.index.year >= 2020)):
            for high in (False, True):
                condition = weekly["us10y"].ge(THRESHOLD) if high else weekly["us10y"].lt(THRESHOLD)
                values = weekly.loc[era & condition, column].dropna()
                rows.append(dict(outcome=key, unit=unit, horizon_weeks=FORWARD_WEEKS, period=period,
                                 regime="High yield" if high else "Lower yield", n=len(values),
                                 median_change=values.median(), p10_change=values.quantile(.1),
                                 p90_change=values.quantile(.9), rising_share_pct=100 * values.gt(0).mean()))
    return pd.DataFrame(rows, columns=["outcome", "unit", "horizon_weeks", "period", "regime", "n",
                                       "median_change", "p10_change", "p90_change", "rising_share_pct"])


# ============================================================================
# SECTION 5. OUTPUT. A single Excel workbook; optional publication images.
# Native Excel charts reference visible source sheets in the same workbook.
# ============================================================================
def methodology(end, manifest):
    return pd.DataFrame([
        ("Scope", "Transparent reconstruction of user-provided Bloomberg screenshots. Exact original tickers, date range, smoothing and trend estimator are not fully disclosed. Published table numbers are not hardcoded."),
        ("Cutoff", f"Data <= {end}. Incomplete final weeks/months excluded. Daily close data cannot recreate the article's intraday snapshot. Historical data are today's downloaded vintage, not archived point-in-time vintages."),
        ("Core instruments", "USGG10YR yield in percent; LUATTRUU Treasury total-return index; SPX price index. SPTR total-return sensitivity is separate. Do not substitute LBUSTRUU (Aggregate) for Treasuries."),
        ("Data flow", "Bloomberg HistoricalDataRequest -> immutable CSV + manifest + API messages -> hash-checked disk read -> derived tables -> charts. Required failures stop; optional failures remain visible."),
        ("Weekly calendar", "Last common actual date for stock, bond and yield inside each Mon-Fri week. W-FRI is a label, source_date is the actual date. No forward fill. A missing week breaks the return calculation and rolling window."),
        ("Monthly history", "A separate MONTHLY Treasury request supports the long YoY chart. Monthly-only history never becomes weekly data. Check source observations by year before comparing with the article."),
        ("Stock-bond correlation", f"Pearson correlation of {CORRELATION_WEEKS} consecutive one-week simple returns. Weekly returns are non-overlapping, but adjacent correlation estimates overlap heavily. Positive includes joint gains as well as losses."),
        ("Sampling sensitivity", f"52/{CORRELATION_WEEKS}/156 weekly windows; SPTR versus SPX; 504 daily observations of overlapping 5 matched-session returns (5-10 elapsed calendar days). Tables compare own histories and common weekly dates, and show actual daily-window calendar spans; 504 matched sessions need not equal two years."),
        ("Yield buckets", f"[lower, upper), including {THRESHOLD}% in the high regime. Counts require both yield and correlation. Positive/negative/zero percentages share the same denominator. Runs and calendar years are descriptive, not independent sample sizes."),
        ("Trending mean", "Treasury YoY = 100*(month-end total-return index / index 12 calendar months earlier - 1). Full-sample OLS on elapsed years; bands = fitted trend +/- 1 or 2 * sqrt(SSE/(n-2)). This reconstruction is an assumption, not Bloomberg's disclosed formula."),
        ("Band interpretation", f"Residual SD bands are historical dispersion, not forecast confidence intervals or fair value. Full-sample trend uses future observations relative to historical chart dates. A separate prior-only fit starts after {MIN_TREND_MONTHS} valid prior YoY observations."),
        ("Prior-only z", "At t, fit linear trend to valid observations strictly before t, then divide actual-minus-predicted YoY by that training sample's residual SD. No claim that z is normally distributed or that the return level mean-reverts predictably."),
        ("Forward returns", f"At t, compound returns from t+1 through t+{FORWARD_WEEKS}; every intervening week is required. Joint loss means both cumulative returns <0. Last {FORWARD_WEEKS} weeks have no matured outcome. SPX results exclude dividends."),
        ("60/40 diagnostic", "Weekly rebalanced 60% SPX price return + 40% Treasury total return, then compounded. This is a price-return diagnostic, not a total-return strategy backtest; fees and trading costs omitted."),
        ("Event study", f"First week >= {THRESHOLD}% after 4 consecutive below-threshold weeks. Event windows must mature and start more than {FORWARD_WEEKS} weeks apart. Endpoints of optional index changes must match the core source dates. Small episode counts limit inference."),
        ("Bootstrap", "Circular paired moving blocks of 26/52/104 calendar weeks; missing rows stay on the grid. Outcome difference = high-yield mean minus lower-yield mean. Each regime needs >=20 observations. CI withheld if fewer than 80% (and 100) valid replications. Exploratory, no multiple-testing correction, not proof of causality/stationarity."),
        ("Non-overlap check", f"Evaluate every one of the {FORWARD_WEEKS} possible fixed-phase grids of {FORWARD_WEEKS}-week outcomes; report all phases. Non-overlapping return windows can still be serially dependent."),
        ("Volatility / liquidity", "Plot raw weekly matched-date values and explicit 13-week averages separately. Original screenshot smoother unknown. Higher GVLQUSD means poorer liquidity; MOVE and VIX are different implied-volatility measures. Correlation of changes is separate from level correlation."),
        ("Forward transmission test", f"For MOVE, VIX, HY OAS and liquidity, compare the next {FORWARD_WEEKS}-week change across yield regimes and eras. Each outcome uses its own valid endpoint pairs; sample counts differ. P10/P90 are outcome dispersion, not confidence intervals. No causal chain is assumed."),
        ("Credit units", "LF98OAS PX_LAST raw is in percent; derived hy_oas is multiplied by 100 for bp. Raw CSV/Excel retain the unscaled API values. Check DES/FLDS and entitlements when changing ticker or field."),
        ("Commodity scope", "BCOM is a global commodity benchmark used as a US inflation context variable. Its rising level does not establish commodity breadth; no non-US macro releases are included."),
        ("Research limits", "5.25% is a hypothesis from the article, not an invariant economic boundary. Era mix, trend inflation, duration changes and shock composition may confound the relation. A shift in return correlation alone does not prove MOVE, VIX or spreads will rise."),
        ("Not replicated", "Stock-level implied correlation/dispersion, exact FOMC reaction and commodity-sector breadth require additional definitions/data. No fabricated proxy is inserted."),
        ("Next extension", "To distinguish discount-rate shocks from growth news, add verified real-yield, breakeven, term-premium and timestamped macro-surprise histories. Use pre-specified out-of-period tests; revised macro data require vintage control."),
        ("Snapshot", f"Retrieved {manifest['retrieved_at']}; raw SHA256 {manifest['sha256']}"),
        *[("Source: " + key, value) for key, value in SOURCES.items()],
    ], columns=["topic", "method"])


def create_charts(weekly, trend, tables, pairs):
    # Imported only for analysis/check, so fetch does not require Matplotlib.
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10, "axes.spines.top": False,
                         "axes.spines.right": False, "axes.labelcolor": "#516071", "text.color": "#162c43",
                         "axes.edgecolor": "#ccd3da", "axes.titleweight": "bold", "svg.fonttype": "none",
                         "savefig.facecolor": "white"})
    navy, orange, teal, gray = "#173f65", "#dc851f", "#318a88", "#a2aab3"
    charts = []

    def new(title, subtitle, ylabel, size=(11.5, 5.4)):
        figure, ax = plt.subplots(figsize=size)
        figure.subplots_adjust(left=.08, right=.91, bottom=.20, top=.79)
        figure.text(.08, .94, title, fontsize=16, weight="bold")
        figure.text(.08, .885, subtitle, fontsize=9, color="#596978")
        ax.set_ylabel(ylabel)
        ax.grid(axis="y", color="#e5e9ed", linewidth=.7)
        ax.set_axisbelow(True)
        return figure, ax

    def add(key, title, figure, ax, frame, columns, ylabel, note, kind="line", second_axis=None):
        figure.text(.08, .04, "Source: Bloomberg; own calculations. " + note, fontsize=8, color="#657383", wrap=True)
        handles, labels = ax.get_legend_handles_labels()
        if handles:
            ax.legend(handles, labels, loc="lower left", bbox_to_anchor=(0, 1.01), borderaxespad=0,
                      frameon=False, fontsize=8, ncol=min(6, len(handles)))
        charts.append(dict(key=key, title=title, figure=figure, data=frame[columns].copy(),
                           columns=columns, ylabel=ylabel, note=note, kind=kind, second_axis=second_axis))

    title = "Treasury returns relative to their historical trend"
    fig, ax = new(title, "12-month total return · full-sample linear trend and residual SD bands", "12M return (%)")
    ax.fill_between(trend.index, trend["minus_2sd"], trend["plus_2sd"], color=teal, alpha=.055)
    for name, color, label in (("minus_2sd", gray, "−2 SD"), ("minus_1sd", gray, "−1 SD"),
                                ("plus_1sd", teal, "+1 SD"), ("plus_2sd", teal, "+2 SD"),
                                ("trend_pct", navy, "Trending mean"), ("yoy_pct", orange, "Treasury YoY")):
        ax.plot(trend.index, trend[name], color=color, label=label, lw=1.9 if name == "yoy_pct" else 1)
    add("01_treasury_trend", title, fig, ax, trend,
        ["yoy_pct", "trend_pct", "minus_1sd", "plus_1sd", "minus_2sd", "plus_2sd"], "%",
        "Trend estimator is a disclosed reconstruction; bands are not forecast intervals.")

    title = "Where is stock-bond correlation positive?"
    bucket = tables["yield_buckets"].iloc[1:-1].set_index("bucket")
    fig, ax = new(title, f"10Y yield buckets · {CORRELATION_WEEKS}-week return correlation · labels show positive share and sample count", "Share of observations (%)", (11.5, 6.2))
    x = np.arange(len(bucket))
    ax.bar(x, bucket["positive_pct"], color=navy, label="Positive correlation")
    ax.bar(x, bucket["negative_pct"], bottom=bucket["positive_pct"], color="#dfe5eb", label="Negative correlation")
    for i, (_, row) in enumerate(bucket.iterrows()):
        label = f"{row['positive_pct']:.0f}%\nn={int(row['n'])}" if row["n"] else "n=0"
        ax.text(i, 103, label, ha="center", va="bottom", fontsize=7)
    ax.set_xticks(x, bucket.index, rotation=50, ha="right", fontsize=8)
    ax.set_ylim(0, 123)
    ax.set_yticks([0, 25, 50, 75, 100])
    ax.legend(loc="upper left", fontsize=8, frameon=False)
    add("02_yield_buckets", title, fig, ax, bucket, ["positive_pct", "negative_pct", "zero_pct"], "%",
        "Buckets include lower bound. Overlapping correlations are not independent observations.", "bar")

    title = "Yield level and the stock-bond regime"
    fig, ax = new(title, f"Same-date weekly observations · 52-, {CORRELATION_WEEKS}- and 156-week sensitivity", "Return correlation")
    correlation_series = dict.fromkeys(["corr_52w", "corr_156w", CORRELATION_COLUMN])
    for name in correlation_series:
        color = navy if name == CORRELATION_COLUMN else (gray if name == "corr_52w" else teal)
        label = name.removeprefix("corr_").removesuffix("w") + " weeks"
        ax.plot(weekly.index, weekly[name], color=color, lw=1.3, label=label)
    ax.axhline(0, color="#8695a5", lw=.7)
    ax.set_ylim(-1, 1)
    ax.fill_between(weekly.index, -1, 1, where=weekly["us10y"].ge(THRESHOLD), color=orange, alpha=.10)
    add("03_correlation_history", title, fig, ax, weekly, list(correlation_series), "Correlation",
        f"Shaded weeks: 10Y >= {THRESHOLD}%. Missing windows remain blank.")

    # Raw levels are plotted separately from averages, so smoothing is visible.
    for key, pair in pairs.items():
        first, second = pair.columns
        labels = {"move": "MOVE", "liquidity": "Treasury liquidity (higher = worse)", "vix": "VIX", "hy_oas": "HY OAS (bp)"}
        title = f"{labels[first]} and {labels[second]}"
        fig, ax = new(title, "Raw weekly observations in pale lines · explicit 13-week averages in solid lines", labels[first])
        other = ax.twinx()
        other.spines["top"].set_visible(False)
        other.spines["right"].set_visible(True)
        averaged = pair.rolling(13).mean().rename(columns=lambda c: c + "_13w")
        for axis, column, color in ((ax, first, navy), (other, second, orange)):
            axis.plot(pair.index, pair[column], color=color, alpha=.22, lw=.6)
            axis.plot(pair.index, averaged[column + "_13w"], color=color, label=labels[column], lw=1.8)
        other.set_ylabel(labels[second], color=orange)
        other.legend(loc="lower right", bbox_to_anchor=(1, 1.01), borderaxespad=0, frameon=False, fontsize=8)
        add("04_" + key, title, fig, ax, averaged, list(averaged.columns), labels[first],
            "Dual axes; exact source dates match. Co-movement does not identify causation.", second_axis=labels[second])
    if "bcom" in weekly:
        title = "Commodity prices as inflation context"
        frame = weekly.loc[weekly.index.year >= 2012]
        fig, ax = new(title, "Bloomberg Commodity Index · weekly level", "BCOM index")
        ax.plot(frame.index, frame["bcom"], color=orange, lw=1.8)
        add("05_commodities", title, fig, ax, frame, ["bcom"], "Index",
            "A higher aggregate index does not by itself establish sector breadth.")

    title = f"Does the {THRESHOLD:g}% result survive different periods?"
    stats = tables["period_regimes"]
    fig, ax = new(title, f"Positive {CORRELATION_WEEKS}-week stock-bond correlation share · high versus lower 10Y yields", "Positive correlation (%)")
    names = list(stats["period"].unique())
    for shift, regime, color in ((-.18, f"< {THRESHOLD:g}%", gray), (.18, f">= {THRESHOLD:g}%", navy)):
        values = stats.loc[stats["regime"].eq(regime)].set_index("period").reindex(names)
        x = np.arange(len(names)) + shift
        ax.bar(x, values["positive_correlation_pct"], width=.34, label=regime, color=color)
        for i, (_, row) in enumerate(values.iterrows()):
            ax.text(x[i], (row["positive_correlation_pct"] if pd.notna(row["positive_correlation_pct"]) else 0) + 3,
                    f"n={int(row['correlation_n'])}", ha="center", fontsize=8)
    ax.set_xticks(np.arange(len(names)), names)
    ax.set_ylim(0, 118)
    period_plot = stats.pivot(index="period", columns="regime", values="positive_correlation_pct").reindex(names)
    add("06_period_stability", title, fig, ax, period_plot, list(period_plot.columns), "%",
        "No observations in a regime means no estimate, not a zero probability.", "bar")

    title = "Is there a unique yield threshold?"
    th = tables["thresholds"].set_index("threshold_pct")
    fig, ax = new(title, "A pre-specified 4.00–6.50% grid · sample counts are in Excel", "Positive correlation (%)")
    ax.plot(th.index, th["high_positive_pct"], color=navy, label="At / above threshold", marker="o", ms=3)
    ax.plot(th.index, th["low_positive_pct"], color=gray, label="Below threshold", marker="o", ms=3)
    ax.axvline(THRESHOLD, color=orange, ls="--", lw=1)
    ax.set_xlabel("10Y yield threshold (%)")
    ax.set_ylim(0, 105)
    add("07_threshold_sensitivity", title, fig, ax, th, ["high_positive_pct", "low_positive_pct"], "%",
        "Exploratory threshold sensitivity; not an optimized trading rule.")

    title = "Do stocks and bonds subsequently lose together?"
    r = tables["yield_regimes"].copy()
    r["state"] = r["yield_level"] + " yield\n" + r["direction"]
    r = r.set_index("state")
    fig, ax = new(title, f"Joint negative cumulative return over the following {FORWARD_WEEKS} weeks", "Joint loss frequency (%)")
    ax.bar(np.arange(len(r)), r["joint_loss_pct"], color=[gray, teal, navy, orange])
    for i, (_, row) in enumerate(r.iterrows()):
        y = row["joint_loss_pct"] if pd.notna(row["joint_loss_pct"]) else 0
        ax.text(i, y + 2, f"n={int(row['n'])}", ha="center", fontsize=9)
    ax.set_xticks(np.arange(len(r)), r.index)
    ax.set_ylim(0, 105)
    add("08_forward_joint_losses", title, fig, ax, r, ["joint_loss_pct"], "%",
        f"Only matured outcomes; overlapping {FORWARD_WEEKS}-week windows. Level and direction known at t.", "bar")

    events = tables["events"]
    if not events.empty:
        title = "What followed sustained upward threshold crossings?"
        ev = events.set_index("date")[["forward_stock_pct", "forward_bond_pct"]]
        fig, ax = new(title, f"First crossing after 4 below-threshold weeks · non-overlapping {FORWARD_WEEKS}-week event windows", f"Next {FORWARD_WEEKS}-week return (%)")
        x = np.arange(len(ev))
        ax.bar(x - .18, ev["forward_stock_pct"], .34, color=navy, label="S&P 500 price return")
        ax.bar(x + .18, ev["forward_bond_pct"], .34, color=orange, label="Treasury total return")
        ax.set_xticks(x, [str(d.date()) for d in ev.index], rotation=40, ha="right", fontsize=8)
        ax.axhline(0, color=gray, lw=.7)
        add("09_crossing_events", title, fig, ax, ev, list(ev.columns), "%",
            "Case studies; few historical episodes cannot establish a stable causal effect.", "bar")
    return charts


def excel_export(path, raw, tables, charts):
    """Single workbook, native editable charts, visible chart source sheets.

    This creates a new version each run. Author chart links are platform-specific;
    this file does not overwrite an existing linked publication workbook.
    """
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.chart.axis import DateAxis
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    workbook = Workbook(write_only=True)

    def clean(value):
        if pd.isna(value):
            return None
        if isinstance(value, pd.Timestamp):
            return value.to_pydatetime()
        if isinstance(value, np.generic):
            value = value.item()
        # Do not allow imported metadata to become an Excel formula.
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
            return "'" + value
        return value

    def sheet(name, frame, keep_index=False):
        frame = frame.reset_index(names="date_or_bucket") if keep_index else frame
        ws = workbook.create_sheet(name)
        ws.freeze_panes = "B2"
        for i, col in enumerate(frame.columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = 25 if i > 1 else 30
        if name == "Methods":
            ws.column_dimensions["B"].width = 115
        if name == "Findings":
            ws.column_dimensions["A"].width = 120
        headers = []
        for label in frame.columns:
            cell = WriteOnlyCell(ws, str(label))
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="173F65")
            headers.append(cell)
        ws.append(headers)
        for row_number, row in enumerate(frame.itertuples(index=False, name=None), 2):
            cells = []
            for value in row:
                value = clean(value)
                cell = WriteOnlyCell(ws, value)
                if isinstance(value, float):
                    cell.number_format = "0.00;-0.00;0.00"
                if name in ("Methods", "Findings"):
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                cells.append(cell)
            if name in ("Methods", "Findings"):
                ws.row_dimensions[row_number].height = max(30, math.ceil(max(len(str(v)) for v in row) / 110)*16)
            ws.append(cells)
        ws.auto_filter.ref = f"A1:{get_column_letter(len(frame.columns))}{len(frame)+1}"
        return ws

    sheet("Findings", tables["Findings"])
    display = workbook.create_sheet("Charts")
    display.sheet_view.showGridLines = False
    display.append(["Treasury regime study — editable charts; chart data on C01, C02, ..."])
    display.column_dimensions["A"].width = 24
    sheet("Methods", tables["methods"])
    for name, frame in tables.items():
        if name not in ("Findings", "methods"):
            sheet(name[:31], frame, isinstance(frame.index, pd.DatetimeIndex))
    for number, spec in enumerate(charts, 1):
        chart_data = spec["data"].copy()
        is_date = isinstance(chart_data.index, pd.DatetimeIndex)
        if spec["kind"] == "bar" and is_date:
            chart_data.index = chart_data.index.strftime("%Y-%m-%d")
        source = sheet(f"C{number:02d}_{spec['key'][:22]}", chart_data, True)
        if not len(spec["data"]):
            continue
        chart = BarChart() if spec["kind"] == "bar" else LineChart()
        chart.title, chart.y_axis.title = spec["title"], spec["ylabel"]
        chart.width, chart.height = 27, 12
        chart.style = 13
        chart.display_blanks = "gap"
        chart.legend.position = "b"
        if spec["key"] == "02_yield_buckets":
            chart.grouping, chart.overlap = "stacked", 100
            chart.y_axis.scaling.min, chart.y_axis.scaling.max = 0, 100
        if spec["kind"] == "line" and is_date:
            chart.x_axis = DateAxis(axId=10, crossAx=100, numFmt="yyyy", majorTimeUnit="years")
            chart.y_axis.crossAx = 10
        n = len(spec["data"]) + 1
        if spec["second_axis"]:
            chart.add_data(Reference(source, min_col=2, max_col=2, min_row=1, max_row=n), titles_from_data=True)
            other = LineChart()
            other.add_data(Reference(source, min_col=3, max_col=3, min_row=1, max_row=n), titles_from_data=True)
            other.y_axis.axId = 200
            other.y_axis.title = spec["second_axis"]
            other.y_axis.crosses = "max"
            if is_date:
                other.x_axis = DateAxis(axId=10, crossAx=200, numFmt="yyyy", majorTimeUnit="years")
            other.set_categories(Reference(source, min_col=1, min_row=2, max_row=n))
            chart += other
        else:
            chart.add_data(Reference(source, min_col=2, max_col=len(spec["data"].columns)+1, min_row=1, max_row=n), titles_from_data=True)
        chart.set_categories(Reference(source, min_col=1, min_row=2, max_row=n))
        display.add_chart(chart, f"A{3 + (number-1)*25}")
    sheet("Raw_data", raw)
    workbook.save(path)


def dynamic_findings(weekly, trend, tables):
    """Short factual text computed after each run; never a fixed market call."""
    last = weekly.dropna(subset=[CORRELATION_COLUMN, "us10y"]).iloc[-1]
    findings = [f"Latest usable correlation: {last[CORRELATION_COLUMN]:+.2f}; 10Y {last['us10y']:.2f}% on {last['source_date'].date()}."]
    stats = tables["period_regimes"].query("period == 'All'")
    hi = stats.loc[stats["regime"].eq(f">= {THRESHOLD:g}%")].iloc[0]
    lo = stats.loc[stats["regime"].eq(f"< {THRESHOLD:g}%")].iloc[0]
    if hi["correlation_n"] and lo["correlation_n"]:
        findings.append(f"Positive correlation share: {hi['positive_correlation_pct']:.0f}% above / at {THRESHOLD}% vs {lo['positive_correlation_pct']:.0f}% below; {int(hi['correlation_n'])} / {int(lo['correlation_n'])} overlapping weekly observations.")
    findings.append(f"Eligible upward-crossing episodes: {len(tables['events'])}. These have fully observed {FORWARD_WEEKS}-week outcomes.")
    latest_z = trend["prior_z"].dropna()
    if len(latest_z):
        findings.append(f"Treasury 12M return versus its prior-only trend: {latest_z.iloc[-1]:+.2f} residual SD on {latest_z.index[-1].date()}.")
    if hi["forward_n"] and lo["forward_n"]:
        findings.append(f"Next-{FORWARD_WEEKS}W joint loss frequency: {hi['joint_loss_pct']:.0f}% in the high-yield regime vs {lo['joint_loss_pct']:.0f}% below. These are separate from correlation-sign shares.")
    return findings


def export(raw, manifest, end, weekly, trend, coverage, daily_sensitivity,
           repetitions, root=ROOT, images=False):
    """Default output: one Excel plus provenance JSON. Images are opt-in."""
    folder = root / "output" / timestamp()
    folder.mkdir(parents=True)
    tables = dict(weekly=weekly, monthly_trend=trend, source_coverage=coverage,
                  daily_sensitivity=daily_sensitivity, yield_buckets=bucket_table(weekly),
                  period_regimes=conditional_summary(weekly), **deeper_research(weekly, daily_sensitivity, repetitions))
    tables["transmission"], pairs = transmission_tables(raw, end)
    tables["forward_channels"] = forward_channels(weekly)
    tables["methods"] = methodology(end, manifest)
    findings = dynamic_findings(weekly, trend, tables)
    for key, error in manifest.get("errors", {}).items():
        findings.append(f"Unavailable optional history: {key}. {error}")
    last_close = weekly.dropna(subset=CORE)["source_date"].iloc[-1]
    last_correlation = weekly.dropna(subset=[CORRELATION_COLUMN])["source_date"].iloc[-1]
    if (pd.Timestamp(end) - last_close).days > 10:
        findings.append(f"Core data are stale: latest common close {last_close.date()}.")
    if last_close != last_correlation:
        findings.append("Latest correlation is older than the latest prices because of missing observations in its window.")
    tables = {"Findings": pd.DataFrame({"finding": findings}), **tables}
    charts = create_charts(weekly, trend, tables, pairs)
    import matplotlib.pyplot as plt
    try:
        if images:
            (folder / "charts").mkdir()
            for spec in charts:
                for extension in ("svg", "png", "pdf"):
                    spec["figure"].savefig(folder / "charts" / f"{spec['key']}.{extension}", dpi=300)
        print("Writing workbook with native charts...", flush=True)
        excel_export(folder / "research.xlsx", raw.loc[raw["date"].le(pd.Timestamp(end))], tables, charts)
        save_json(folder / "run_manifest.json", dict(version=VERSION, asof=str(end), raw_sha256=manifest["sha256"],
                  script_sha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
                  libraries=dict(pandas=pd.__version__, numpy=np.__version__),
                  parameters=dict(threshold=THRESHOLD, correlation_weeks=CORRELATION_WEEKS,
                                  forward_weeks=FORWARD_WEEKS, prior_trend_min_months=MIN_TREND_MONTHS,
                                  bootstrap_repetitions=repetitions),
                  artifacts={str(p.relative_to(folder)): hashlib.sha256(p.read_bytes()).hexdigest()
                             for p in folder.rglob("*") if p.is_file()}))
        # Publish only after every artifact has been written successfully.
        save_json(root / "output" / "latest.json", {"folder": folder.name})
    finally:
        for spec in charts:
            plt.close(spec["figure"])
    print("Excel:", folder / "research.xlsx")
    return folder


# ============================================================================
# SECTION 6. LOCAL CHECKS / COMMAND LINE. Test data never enter real outputs.
# ============================================================================
def self_check():
    """Small exact-value checks; no Bloomberg session, no saved fake history."""
    idx = pd.date_range("2020-01-03", periods=10, freq="W-FRI")
    returns = pd.Series([.01] * 10, index=idx)
    forward = forward_compound(returns, 3)
    assert np.isclose(forward.iloc[0], 1.01 ** 3 - 1)
    assert forward.iloc[-3:].isna().all()
    returns.iloc[2] = np.nan
    assert pd.isna(forward_compound(returns, 3).iloc[0])
    daily = pd.DataFrame({"a": [100., 110., 121.], "b": [100., np.nan, 121.]},
                         index=pd.to_datetime(["2020-01-03", "2020-01-10", "2020-01-17"]))
    sampled = weekly_sample(daily, date(2020, 1, 17))
    assert sampled.loc["2020-01-10", ["a", "b"]].isna().all()
    assert pd.isna(sampled["a"].pct_change(fill_method=None).iloc[-1])
    assert weekly_sample(daily, date(2020, 1, 16)).index.max() == pd.Timestamp("2020-01-10")
    const = pd.Series([1.] * 110)
    assert valid_correlation(const, const, 104).isna().all()
    f = pd.DataFrame({"us10y": [5., 5.25, 5.49, 5.5, 7.5], CORRELATION_COLUMN: [-.1, .2, 0, -.3, .4]}, index=idx[:5])
    buckets = bucket_table(f)
    row = buckets.loc[buckets["bucket"].eq("5.25–5.5%")].iloc[0]
    assert row["n"] == 2 and row["positive_pct"] == 50 and row["zero_pct"] == 50
    assert buckets["n"].sum() == 5
    months = pd.date_range("1980-01-31", periods=180, freq=pd.offsets.MonthEnd())
    history = pd.Series(100 * 1.004 ** np.arange(180), index=months)
    trend = trend_analysis(history)
    assert np.allclose(trend["yoy_pct"].dropna(), (1.004 ** 12 - 1) * 100)
    changed = history.copy()
    changed.iloc[-12:] *= 1.2
    pd.testing.assert_frame_equal(trend.iloc[:-12][["prior_z", "prior_trend_pct"]],
                                  trend_analysis(changed).iloc[:-12][["prior_z", "prior_trend_pct"]])
    print("PASS: forward timing, maturity, missing weeks, partial weeks, constant-series correlation, bucket edges, calendar YoY, prior-only trend.")


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("command", choices=["fetch", "analyze", "run", "check"])
    parser.add_argument("--start", type=date.fromisoformat, help=f"Fetch start only (default {START}); analyze uses the saved history")
    # Local company-machine calendar date; no Windows timezone database dependency.
    parser.add_argument("--asof", type=date.fromisoformat,
                        default=date.today() - timedelta(days=1))
    parser.add_argument("--snapshot", type=Path, help="Read a specific committed raw/<timestamp> folder")
    parser.add_argument("--bootstrap", type=int, default=BOOTSTRAPS, help="Bootstrap repetitions, >= 100")
    parser.add_argument("--images", action="store_true", help="Also export publication charts as SVG, PNG and PDF")
    args = parser.parse_args(argv)
    if (args.start is not None and args.start > args.asof) or args.bootstrap < 100:
        parser.error("Require --start <= --asof and --bootstrap >= 100")
    if args.command in ("analyze", "check") and args.start is not None:
        parser.error("--start is for fetch/run only; analyze uses all saved history up to --asof")
    if args.command == "fetch" and args.images:
        parser.error("--images is for analyze/run only")
    if not (isinstance(CORRELATION_WEEKS, int) and CORRELATION_WEEKS >= 8 and
            isinstance(FORWARD_WEEKS, int) and 1 <= FORWARD_WEEKS <= 104 and
            math.isfinite(THRESHOLD) and MIN_TREND_MONTHS >= 24):
        parser.error("Invalid research settings in section 1")
    if args.snapshot is not None and args.command != "analyze":
        parser.error("--snapshot is only used with analyze")
    if args.command == "check":
        self_check()
        return
    if args.command in ("fetch", "run"):
        fetch(args.start or date.fromisoformat(START), args.asof)
    if args.command in ("analyze", "run"):
        raw, manifest = read_snapshot(snapshot=args.snapshot)
        end = min(args.asof, date.fromisoformat(manifest["end"]))
        if args.asof > end:
            print(f"Analysis capped at snapshot end {end}; fetch again for newer observations.")
        weekly, trend, daily, coverage = calculate(raw, end)
        export(raw, manifest, end, weekly, trend, coverage, daily, args.bootstrap, images=args.images)


if __name__ == "__main__":
    try:
        main()
    except (ValueError, RuntimeError, OSError) as error:
        print("ERROR:", error, file=sys.stderr)
        sys.exit(1)
