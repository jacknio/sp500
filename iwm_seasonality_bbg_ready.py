# -*- coding: utf-8 -*-
"""IWM-only seasonality: standalone Bloomberg research with chart-ready Excel data and separate figure files.

QUICK START (Python 3.9+, Bloomberg Terminal logged in)
    python iwm_seasonality_bbg_ready.py
    python iwm_seasonality_bbg_ready.py --backend xbbg
    python iwm_seasonality_bbg_ready.py --backend pdblp --as-of 2026-10-25
    python iwm_seasonality_bbg_ready.py --backend blpapi

OFFLINE / EXISTING COMPANY DATA
    python iwm_seasonality_bbg_ready.py --csv daily.csv
    CSV: Date,Stock,Benchmark,US10y  (only Date and Stock are required).
    Values must be index LEVELS, not daily returns. Ticker column names also work.
    The CSV return basis is whatever you supplied: this script cannot infer it.

DEPENDENCIES
    numpy, pandas, matplotlib, openpyxl are all required for complete output.
    Bloomberg: an existing working xbbg / pdblp / blpapi installation and access.
    Example: python -m pip install numpy pandas matplotlib openpyxl
    Do not install a new Bloomberg wrapper if your company already has one.
    API docs: https://xbbg.org/python/quickstart
              https://matthewgilbert.github.io/pdblp/tutorial.html
              https://bloomberg.github.io/blpapi-docs/python/

DEFAULTS / INTERPRETATION
    IWM US Equity / PX_LAST is split-adjusted PRICE data: cash dividends are excluded.
    IWM began in 2000; do not claim a 30-year IWM historical sample.
    No automatic ticker guessing, index/ETF splicing, dividend adjustment,
    or missing-history backfill. IWM returns are not total returns.
    With --as-of 2026-10-25, the 25 complete years are 2001-2025.
    Requests start in December 2000 for the initial month-end baseline.
    Missing/short history fails by default. --allow-short-history explicitly
    permits valid available years, always showing the actual sample size.

OUTPUT (default: output folder next to this file)
    report.html: standalone offline report with embedded charts
    01..05 PNG + SVG charts, charts.pdf, raw_daily.csv, coverage.csv,
    monthly_returns.csv, nov_dec_by_year.csv, summary.json, methodology.txt,
    results.xlsx with numeric data tabs and formula audit (NO embedded Excel charts).

TEST / PREVIEW (NO BLOOMBERG NEEDED)
    python iwm_seasonality_bbg_ready.py --self-test
    python iwm_seasonality_bbg_ready.py --demo --out demo_output
    Demo outputs are prominently watermarked; they are NOT market evidence.

If your company uses a proprietary pull function, replace pull_bloomberg() only.
Return a pandas DataFrame with a DatetimeIndex and Stock/Benchmark/US10y columns.
The rest of the analysis does not depend on any Bloomberg wrapper.
"""

# ------------------------- 1. SETTINGS ---------------------------------------
YEARS = 25
AS_OF_DATE = None                  # e.g. "2026-10-25"; None = computer date
START_DATE = None                  # None = automatic baseline month before sample
OUTPUT_DIR = None                  # e.g. r"H:\research\iwm_seasonality"
BACKEND = "auto"                  # auto, xbbg, pdblp, blpapi
TICKERS = {
    "Stock": "IWM US Equity",
    "Benchmark": "SPX Index",      # off: IWM-only mode
    "US10y": "USGG10YR Index",      # off: seasonality does not need yields
    "Bond": "LUATTRUU Index",      # off: seasonality does not need bonds
}
INCLUDE_BENCHMARK = False
INCLUDE_RATES = False
INCLUDE_BOND = False
FIELD = "PX_LAST"
RETURN_BASIS = "Price return (cash dividends excluded)"
ALLOW_SHORT_HISTORY = False
BBG_HOST = "localhost"
BBG_PORT = 8194
REQUEST_TIMEOUT_SECONDS = 120
MIN_OBSERVATIONS_PER_MONTH = 12     # rejects monthly/sparse data passed as daily
MAX_MONTH_END_STALENESS_DAYS = 7
MAX_GAP_CALENDAR_DAYS = 10          # a data-quality screen, not an exchange calendar
BOOTSTRAP_DRAWS = 10000
RANDOM_SEED = 20261025
ROUND_TRIP_COST_BPS = 0.0          # optional Nov-Dec execution cost; charts remain gross
YIELD_SCALE = 1.0                  # USGG10YR normally uses 4.5 to mean 4.5%, not 0.045
# ---------------------------------------------------------------------------

import argparse
import base64
import calendar
import hashlib
import html
import importlib.util
import json
import logging
import math
import os
from pathlib import Path
import sys
import tempfile
import time
import warnings

try:
    import numpy as np
    import pandas as pd
except ImportError as exc:
    raise SystemExit("Install numpy and pandas in this Python environment first: " + str(exc))

LOG = logging.getLogger("iwm_seasonality")
GREEN, INK, MUTED, RED, PAPER = "#16684d", "#172823", "#7c8982", "#ac594a", "#f7f6f2"


def active_tickers():
    keys = ["Stock"]
    if INCLUDE_BENCHMARK: keys.append("Benchmark")
    if INCLUDE_RATES: keys.append("US10y")
    if INCLUDE_BOND: keys.append("Bond")
    return {key: TICKERS[key] for key in keys}


def settings(args):
    asof = pd.Timestamp(args.as_of or AS_OF_DATE or pd.Timestamp.today().date()).normalize()
    if pd.isna(asof): raise ValueError("Invalid AS_OF_DATE")
    years = args.years
    if years < 2: raise ValueError("Use at least two complete calendar years")
    last_year = asof.year - 1
    first_year = last_year - years + 1
    start = pd.Timestamp(START_DATE) if START_DATE else pd.Timestamp(first_year - 1, 12, 1)
    end = pd.Timestamp(last_year, 12, 31)
    if start >= end: raise ValueError("START_DATE must precede the final sample date")
    if not 0 <= ROUND_TRIP_COST_BPS < 20000: raise ValueError("Invalid transaction cost")
    return {"as_of": str(asof.date()), "years": years, "first_year": first_year,
            "last_year": last_year, "start": start, "end": end,
            "allow_short": args.allow_short_history or ALLOW_SHORT_HISTORY,
            "tickers": active_tickers(), "field": FIELD, "basis": RETURN_BASIS,
            "demo": args.demo, "source": "", "warnings": []}


def clean_daily(frame, aliases):
    """Preserve NaNs: no forward-filling of source observations."""
    frame = frame.copy()
    if not isinstance(frame.index, pd.DatetimeIndex):
        if pd.api.types.is_numeric_dtype(frame.index):
            raise ValueError("Daily data need explicit dates, not numeric row numbers")
        frame.index = pd.to_datetime(frame.index, errors="raise")
    if frame.index.tz is not None: frame.index = frame.index.tz_localize(None)
    frame.index = frame.index.normalize()
    if frame.index.has_duplicates: raise ValueError("Duplicate dates in daily data; resolve before analysis")
    rename = {str(k).strip().lower(): k for k in aliases}
    rename.update({v.strip().lower(): k for k, v in aliases.items()})
    frame.columns = [rename.get(str(c).strip().lower(), str(c).strip()) for c in frame.columns]
    if frame.columns.has_duplicates: raise ValueError("Duplicate logical columns after ticker mapping")
    if "Stock" not in frame: raise ValueError("Missing Stock column (or its configured Bloomberg ticker)")
    frame = frame[[c for c in aliases if c in frame]].apply(pd.to_numeric, errors="coerce")
    for key in [c for c in ["Stock", "Benchmark", "Bond"] if c in frame]:
        if (frame[key].dropna() <= 0).any(): raise ValueError(key + " contains nonpositive index levels")
    if "US10y" in frame: frame["US10y"] *= YIELD_SCALE
    return frame.sort_index().dropna(how="all").rename_axis("Date")


def normalize_bdh(raw, ticker):
    """Accept legacy wide MultiIndex and modern long-form Bloomberg frames."""
    if hasattr(raw, "to_pandas"): raw = raw.to_pandas()
    if isinstance(raw, pd.Series): return raw.rename(ticker)
    if raw is None or len(raw) == 0: raise ValueError("Empty Bloomberg response for " + ticker)
    raw = raw.copy()
    if isinstance(raw.columns, pd.MultiIndex):
        match = [c for c in raw.columns if ticker.upper() in [str(x).upper() for x in c]
                 and FIELD.upper() in [str(x).upper() for x in c]]
        if len(match) != 1: raise ValueError("Cannot identify ticker/field in Bloomberg columns: " + str(raw.columns))
        return raw[match[0]].rename(ticker)
    lower = {str(c).lower(): c for c in raw.columns}
    date_col = next((lower[k] for k in ["date", "time", "datetime"] if k in lower), None)
    if date_col is not None:
        security = next((lower[k] for k in ["ticker", "security"] if k in lower), None)
        if security is not None: raw = raw[raw[security].astype(str).str.upper() == ticker.upper()]
        if "field" in lower: raw = raw[raw[lower["field"]].astype(str).str.upper() == FIELD.upper()]
        value_col = lower.get("value", lower.get(FIELD.lower()))
        if value_col is None: raise ValueError("Cannot identify Bloomberg value column")
        result = pd.Series(raw[value_col].to_numpy(), index=pd.to_datetime(raw[date_col]), name=ticker)
    elif FIELD.lower() in lower:
        result = raw[lower[FIELD.lower()]].rename(ticker)
    elif ticker.lower() in lower:
        result = raw[lower[ticker.lower()]].rename(ticker)
    else:
        raise ValueError("Unsupported Bloomberg response layout: " + str(raw.columns))
    if result.empty: raise ValueError("No values for " + ticker + " / " + FIELD)
    return result


def native_bdh(ticker, start, end):
    """Desktop Bloomberg reference-data request; handles partial responses/errors."""
    import blpapi
    options = blpapi.SessionOptions()
    options.setServerHost(BBG_HOST)
    options.setServerPort(BBG_PORT)
    session = blpapi.Session(options)
    try:
        if not session.start(): raise RuntimeError("Bloomberg session could not start; check Terminal/API connection")
        if not session.openService("//blp/refdata"): raise RuntimeError("Cannot open //blp/refdata")
        request = session.getService("//blp/refdata").createRequest("HistoricalDataRequest")
        request.getElement("securities").appendValue(ticker)
        request.getElement("fields").appendValue(FIELD)
        request.set("startDate", start.strftime("%Y%m%d"))
        request.set("endDate", end.strftime("%Y%m%d"))
        request.set("periodicitySelection", "DAILY")
        request.set("nonTradingDayFillOption", "ACTIVE_DAYS_ONLY")
        request.set("adjustmentNormal", False)
        request.set("adjustmentAbnormal", False)
        request.set("adjustmentSplit", True)
        request.set("adjustmentFollowDPDF", False)
        session.sendRequest(request)
        deadline = time.monotonic() + REQUEST_TIMEOUT_SECONDS
        rows = []
        while time.monotonic() < deadline:
            event = session.nextEvent(1000)
            for msg in event:
                if msg.hasElement("responseError"): raise RuntimeError(str(msg.getElement("responseError")))
                if str(msg.messageType()) in ["RequestFailure", "SessionTerminated"]:
                    raise RuntimeError(str(msg))
                if not msg.hasElement("securityData"): continue
                sec = msg.getElement("securityData")
                if sec.hasElement("securityError"): raise RuntimeError(str(sec.getElement("securityError")))
                if sec.hasElement("fieldExceptions") and sec.getElement("fieldExceptions").numValues():
                    raise RuntimeError(str(sec.getElement("fieldExceptions")))
                data = sec.getElement("fieldData")
                for i in range(data.numValues()):
                    item = data.getValueAsElement(i)
                    if item.hasElement(FIELD) and not item.getElement(FIELD).isNull():
                        rows.append((pd.Timestamp(item.getElementAsDatetime("date")), item.getElementAsFloat(FIELD)))
            if event.eventType() == blpapi.Event.RESPONSE:
                if not rows: raise ValueError("No Bloomberg observations for " + ticker)
                return pd.Series([v for _, v in rows], index=[t for t, _ in rows], name=ticker)
        raise TimeoutError("Bloomberg request timed out: " + ticker)
    finally:
        session.stop()


def pull_bloomberg(config, backend="auto"):
    """Replace THIS function if your firm has its own Bloomberg data loader."""
    if backend == "auto":
        backend = next((name for name in ["xbbg", "pdblp", "blpapi"]
                        if importlib.util.find_spec(name) is not None), None)
    if backend is None:
        raise RuntimeError("No Bloomberg package found. Use your firm's xbbg/pdblp/blpapi environment, or --csv daily.csv")
    config["source"] = "Bloomberg / " + backend
    LOG.info("Backend: %s; request %s to %s", backend, config["start"].date(), config["end"].date())
    collected = {}
    con = None
    try:
        if backend == "pdblp":
            import pdblp
            con = pdblp.BCon(debug=False, port=BBG_PORT, timeout=REQUEST_TIMEOUT_SECONDS * 1000)
            con.start()
        for key, ticker in config["tickers"].items():
            try:
                LOG.info("Pull %s: %s / %s", key, ticker, FIELD)
                if backend == "xbbg":
                    from xbbg import blp
                    raw = blp.bdh(tickers=ticker, flds=FIELD,
                                  start_date=config["start"].strftime("%Y%m%d"),
                                  end_date=config["end"].strftime("%Y%m%d"), Per="D",
                                  adjustmentNormal=False, adjustmentAbnormal=False,
                                  adjustmentSplit=True, adjustmentFollowDPDF=False)
                    series = normalize_bdh(raw, ticker)
                elif backend == "pdblp":
                    raw = con.bdh([ticker], [FIELD], config["start"].strftime("%Y%m%d"),
                                  config["end"].strftime("%Y%m%d"),
                                  elms=[("periodicitySelection", "DAILY"),
                                        ("adjustmentNormal", False), ("adjustmentAbnormal", False),
                                        ("adjustmentSplit", True), ("adjustmentFollowDPDF", False)])
                    series = normalize_bdh(raw, ticker)
                else:
                    series = native_bdh(ticker, config["start"], config["end"])
                series.index = pd.to_datetime(series.index)
                collected[key] = series
            except Exception as exc:
                if key == "Stock": raise RuntimeError("Stock pull failed: " + ticker + " / " + str(exc)) from exc
                message = "Optional " + key + " unavailable: " + str(exc)
                config["warnings"].append(message)
                LOG.warning(message)
    finally:
        if con is not None: con.stop()
    return pd.DataFrame(collected)


def month_data(series):
    series = series.dropna().sort_index()
    grouped = series.groupby(series.index.to_period("M"))
    frame = pd.DataFrame({"level": grouped.last(), "count": grouped.size()})
    dates = pd.Series(series.index, index=series.index)
    frame["actual_date"] = dates.groupby(dates.index.to_period("M")).last()
    frame["stale_days"] = (frame.index.to_timestamp("M") - pd.DatetimeIndex(frame.actual_date)).days
    return frame


def validate_years(stock, config):
    """Require 12 months plus preceding December; never fill a missing month."""
    months = month_data(stock)
    accepted, records = [], []
    for year in range(config["first_year"], config["last_year"] + 1):
        errors = []
        for period in pd.period_range(str(year - 1) + "-12", str(year) + "-12", freq="M"):
            if period not in months.index:
                errors.append(str(period) + " missing")
            else:
                row = months.loc[period]
                if row["count"] < MIN_OBSERVATIONS_PER_MONTH: errors.append(str(period) + " sparse")
                if row.stale_days > MAX_MONTH_END_STALENESS_DAYS: errors.append(str(period) + " stale month-end")
        daily = stock.dropna().loc[str(year - 1) + "-12-01":str(year) + "-12-31"]
        gaps = daily.index.to_series().diff().dt.days
        if len(gaps) and gaps.max() > MAX_GAP_CALENDAR_DAYS:
            errors.append("gap > " + str(MAX_GAP_CALENDAR_DAYS) + " calendar days")
        valid = not errors
        if valid: accepted.append(year)
        records.append({"year": year, "valid": valid, "observations": int(len(stock.loc[str(year)].dropna()))
                        if year in stock.index.year else 0, "reason": "; ".join(errors)})
    return accepted, pd.DataFrame(records), months


def window_return(series, year, first_month, last_month):
    md = month_data(series)
    prior = pd.Period(year=year, month=first_month, freq="M") - 1
    finish = pd.Period(year=year, month=last_month, freq="M")
    required = pd.period_range(prior, finish, freq="M")
    if not required.isin(md.index).all(): return np.nan
    subset = md.loc[required]
    if (subset.stale_days > MAX_MONTH_END_STALENESS_DAYS).any(): return np.nan
    if (subset["count"] < MIN_OBSERVATIONS_PER_MONTH).any(): return np.nan
    return float(md.loc[finish, "level"] / md.loc[prior, "level"] - 1)


def stats(values):
    x = np.asarray(pd.Series(values).dropna(), dtype=float)
    if not len(x): return {"n": 0}
    n = len(x)
    rng = np.random.default_rng(RANDOM_SEED)
    mean_ci = np.quantile(rng.choice(x, (BOOTSTRAP_DRAWS, n), replace=True).mean(axis=1), [.025, .975])
    return {"n": n, "mean": float(x.mean()), "median": float(np.median(x)),
            "win_rate": float((x > 0).mean()), "min": float(x.min()), "max": float(x.max()),
            "std": float(x.std(ddof=1)) if n > 1 else None,
            "mean_bootstrap_95ci": mean_ci.tolist()}


def corr(x, y):
    data = pd.concat([x, y], axis=1).dropna()
    if len(data) < 3 or (data.std() < 1e-12).any(): return np.nan
    return float(data.iloc[:, 0].corr(data.iloc[:, 1]))


def analyze(daily, config, out):
    daily = daily.loc[config["start"]:config["end"]]
    if daily.empty or daily.Stock.dropna().empty: raise ValueError("No stock history in requested dates")
    daily.to_csv(out / "raw_daily.csv", float_format="%.12g")
    years, coverage, md = validate_years(daily.Stock, config)
    coverage.to_csv(out / "coverage.csv", index=False)
    if len(years) != config["years"]:
        bad = coverage.loc[~coverage.valid, ["year", "reason"]].to_string(index=False)
        message = (f"Requested {config['years']} complete years; only {len(years)} passed quality checks.\n"
                   + bad + "\nSee coverage.csv. IWM was launched in 2000; choose a sample with a valid previous-December baseline.\n"
                   "--allow-short-history explicitly permits the valid subset; missing history is NOT filled.")
        if not config["allow_short"]: raise ValueError(message)
        config["warnings"].append(message)
        LOG.warning(message)
    if len(years) < 2: raise ValueError("Fewer than two valid complete years")
    monthly = pd.DataFrame(index=years, columns=range(1, 13), dtype=float)
    annual_paths, seasonal_paths, records = [], [], []
    stock = daily.Stock.dropna()
    half_cost = ROUND_TRIP_COST_BPS / 20000
    for year in years:
        periods = pd.period_range(str(year - 1) + "-12", str(year) + "-12", freq="M")
        levels = md.loc[periods, "level"].to_numpy()
        monthly.loc[year] = levels[1:] / levels[:-1] - 1
        dates = pd.date_range(str(year - 1) + "-12-31", str(year) + "-12-31", freq="D")
        dates = dates[~((dates.month == 2) & (dates.day == 29))]
        series = stock.reindex(stock.index.union(dates)).sort_index().ffill().reindex(dates)
        annual_paths.append((series.to_numpy() / levels[0] - 1))
        sd = pd.date_range(str(year) + "-10-31", str(year) + "-12-31", freq="D")
        base = md.loc[pd.Period(str(year) + "-10", freq="M"), "level"]
        sp = stock.reindex(stock.index.union(sd)).sort_index().ffill().reindex(sd) / base
        seasonal_paths.append(sp.to_numpy() - 1)
        novdec = float(sp.iloc[-1] - 1)
        row = {"year": year, "November": monthly.loc[year, 11], "December": monthly.loc[year, 12],
               "NovDec": novdec, "NovDec_after_cost": (1 - half_cost)**2 * (1 + novdec) - 1,
               "FullYear": levels[-1] / levels[0] - 1,
               "NovDec_max_drawdown": float((sp / sp.cummax() - 1).min())}
        if "Benchmark" in daily:
            row["Benchmark_NovDec"] = window_return(daily.Benchmark, year, 11, 12)
            row["Excess_NovDec"] = novdec - row["Benchmark_NovDec"]
        records.append(row)
    monthly.index.name = "Year"
    annual_paths = np.array(annual_paths)
    seasonal_paths = np.array(seasonal_paths)
    yearly = pd.DataFrame(records).set_index("year")
    monthly.to_csv(out / "monthly_returns.csv")
    yearly.to_csv(out / "nov_dec_by_year.csv")
    pd.DataFrame(annual_paths.T, columns=years).to_csv(out / "annual_paths.csv", index_label="nonleap_day")
    pd.DataFrame(seasonal_paths.T, columns=years).to_csv(out / "nov_dec_paths.csv", index_label="days_after_oct31")
    results = {"NovDec": stats(yearly.NovDec), "November": stats(yearly.November),
               "December": stats(yearly.December), "NovDec_net": stats(yearly.NovDec_after_cost),
               "Last10_NovDec": stats(yearly.loc[yearly.index >= config["last_year"] - 9, "NovDec"])}
    if "Excess_NovDec" in yearly: results["Excess_NovDec"] = stats(yearly.Excess_NovDec)
    results["monthly"] = {str(m): stats(monthly[m]) for m in monthly}
    rates = None
    if "US10y" in daily and len(daily.US10y.dropna()) > 36:
        if daily.US10y.dropna().abs().median() > 25:
            raise ValueError("Yield scale looks wrong. US10y must be percent units (e.g. 4.5), check YIELD_SCALE")
        joint = daily[["Stock", "US10y"]].dropna()
        # Last COMMON observation each month, not independently forward-filled yields.
        common = joint.groupby(joint.index.to_period("M")).last()
        joint_quality = month_data(joint.Stock)
        bad = ((joint_quality["count"] < MIN_OBSERVATIONS_PER_MONTH)
               | (joint_quality.stale_days > MAX_MONTH_END_STALENESS_DAYS))
        common.loc[joint_quality.index[bad], :] = np.nan
        all_months = pd.period_range(str(config["first_year"] - 1) + "-12",
                                     str(config["last_year"]) + "-12", freq="M")
        common = common.reindex(all_months)
        rates = pd.DataFrame({"stock_return": common.Stock / common.Stock.shift(1) - 1,
                              "yield_change_bp": common.US10y.diff() * 100})
        rates = rates.loc[rates.index.year >= config["first_year"]].copy()
        # Preserve empty calendar months so a rolling 36-month window cannot jump a missing year.
        rates.loc[~rates.index.year.isin(years), :] = np.nan
        rates["rolling36"] = rates.stock_return.rolling(36, min_periods=36).corr(rates.yield_change_bp)
        rates.to_csv(out / "rates_monthly.csv", index_label="month")
        complete = rates[["stock_return", "yield_change_bp"]].dropna()
        results["correlation"] = {"n": len(complete), "full": corr(complete.stock_return, complete.yield_change_bp),
                                  "pre2022": corr(complete.loc[:"2021", "stock_return"], complete.loc[:"2021", "yield_change_bp"]),
                                  "since2022": corr(complete.loc["2022":, "stock_return"], complete.loc["2022":, "yield_change_bp"])}
    return {"daily": daily, "monthly": monthly, "yearly": yearly, "years": years,
            "annual_paths": annual_paths, "seasonal_paths": seasonal_paths,
            "coverage": coverage, "stats": results, "rates": rates}


def figures(result, config, out):
    os.environ.setdefault("MPLCONFIGDIR", str(Path(tempfile.gettempdir()) / "iwm_matplotlib"))
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages
    from matplotlib.ticker import PercentFormatter
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10, "figure.facecolor": PAPER,
                         "axes.facecolor": PAPER, "axes.spines.top": False, "axes.spines.right": False,
                         "axes.edgecolor": "#cbd3cb", "text.color": INK, "axes.labelcolor": INK,
                         "grid.color": "#dde2dc", "axes.titleweight": "bold", "svg.fonttype": "path"})
    years, monthly, yearly = result["years"], result["monthly"], result["yearly"]
    label = f"{config['tickers']['Stock']} | {min(years)}–{max(years)} | n={len(years)} | {config['basis']}"
    paths, nd = result["annual_paths"], result["seasonal_paths"]
    names = []
    pdf = PdfPages(out / "charts.pdf")
    def save(fig, name, title, subtitle):
        fig.suptitle(title, x=.07, ha="left", fontsize=17, fontweight="bold", y=.98)
        fig.text(.07,.905,label,fontsize=9,color=MUTED)
        if config["demo"]:
            fig.text(.5,.52,"SYNTHETIC DEMO — NOT BLOOMBERG",ha="center",rotation=15,
                     fontsize=23,color=RED,alpha=.24,weight="bold")
        fig.text(.07,.015,subtitle,fontsize=8,color=MUTED)
        fig.tight_layout(rect=[.01,.055,.99,.87])
        fig.savefig(out / (name + ".png"),dpi=180,bbox_inches="tight")
        fig.savefig(out / (name + ".svg"),bbox_inches="tight")
        pdf.savefig(fig,bbox_inches="tight")
        plt.close(fig)
        names.append((name,title))
    def pct(ax):
        ax.yaxis.set_major_formatter(PercentFormatter(1,decimals=0))
        ax.axhline(0,color=MUTED,lw=.8)
        ax.grid(axis="y",alpha=.45)

    fig,ax=plt.subplots(figsize=(12,5.4));x=np.arange(366)
    low,high=np.quantile(paths,[.25,.75],axis=0)
    ax.fill_between(x,low,high,color=GREEN,alpha=.14,label="Middle 50% of historical years")
    ax.plot(x,paths.mean(axis=0),color=GREEN,lw=2.4,label="Mean")
    ax.plot(x,np.median(paths,axis=0),color=INK,lw=1.7,ls="--",label="Median")
    recent=[i for i,y in enumerate(years) if y>=config["last_year"]-9]
    if recent: ax.plot(x,paths[recent].mean(axis=0),color="#7192b0",lw=1.5,label=f"Latest 10 calendar years (n={len(recent)})")
    ticks=[0]+[(pd.Timestamp(2001,m,calendar.monthrange(2001,m)[1])-pd.Timestamp(2000,12,31)).days for m in range(1,13)]
    ax.set(xticks=ticks,xticklabels=["Start"]+[calendar.month_abbr[m] for m in range(1,13)],xlim=(0,365),ylabel="Cumulative return from prior year-end")
    pct(ax);ax.legend(frameon=False,fontsize=9,loc="upper left")
    save(fig,"01_full_year_seasonality","IWM ETF | the full-year seasonal path","Calendar-day alignment; Feb 29 omitted. Shading shows dispersion, not a confidence interval or forecast.")

    fig,axes=plt.subplots(1,2,figsize=(12,5.1));means=monthly.mean();wins=(monthly>0).mean();x=np.arange(12)
    axes[0].bar(x,means,color=[GREEN if v>=0 else RED for v in means],width=.65)
    for i,v in enumerate(means):axes[0].annotate(f"{v:+.1%}",(i,v),xytext=(0,5 if v>=0 else -14),textcoords="offset points",ha="center",fontsize=8)
    axes[1].bar(x,wins,color=GREEN,alpha=.8,width=.65);axes[1].axhline(.5,color=MUTED,ls="--",lw=1)
    for i,v in enumerate(wins):axes[1].text(i,v+.025,f"{v:.0%}",ha="center",fontsize=8)
    for ax in axes:ax.set(xticks=x,xticklabels=[calendar.month_abbr[m] for m in range(1,13)]);pct(ax)
    axes[0].set_title("Mean monthly return");axes[0].margins(y=.23)
    axes[1].set(title="Frequency of positive months",ylim=(0,1.13))
    save(fig,"02_monthly_seasonality","Which months have historically been strongest?","Each month uses the same accepted years. Historical win rates are not probabilities for the coming year.")

    fig,axes=plt.subplots(1,2,figsize=(12,5.3));x=np.arange(62);low,high=np.quantile(nd,[.25,.75],axis=0)
    axes[0].fill_between(x,low,high,color=GREEN,alpha=.15,label="Middle 50%")
    axes[0].plot(x,nd.mean(axis=0),color=GREEN,lw=2.4,label="Mean")
    axes[0].plot(x,np.median(nd,axis=0),color=INK,ls="--",lw=1.6,label="Median")
    axes[0].set(xticks=[0,15,30,46,61],xticklabels=["Oct 31","Nov 15","Nov 30","Dec 16","Dec 31"],xlim=(0,61),title="Oct-end to Dec-end path")
    axes[0].legend(frameon=False,fontsize=9);pct(axes[0])
    x=np.arange(len(years));axes[1].bar(x,yearly.NovDec,color=[GREEN if v>=0 else RED for v in yearly.NovDec],width=.72)
    axes[1].axhline(yearly.NovDec.mean(),ls="--",lw=1,color=INK)
    axes[1].set(xticks=x[::2],xticklabels=[str(y) for y in years[::2]],title=f"Mean {yearly.NovDec.mean():+.1%} | positive {(yearly.NovDec>0).sum()}/{len(years)}")
    axes[1].tick_params(axis="x",rotation=45,labelsize=8);pct(axes[1])
    save(fig,"03_nov_dec_focus","November–December | distribution, not a guaranteed rally","Entry: October last close. Exit: December last close. Bar chart includes every accepted year; gross of trading costs.")

    heat=monthly.copy();heat[13]=yearly.FullYear
    fig,ax=plt.subplots(figsize=(12,max(8,len(years)*.31+2)))
    image=ax.imshow(heat.to_numpy()*100,cmap="RdYlGn",vmin=-12,vmax=12,aspect="auto")
    for i in range(len(heat)):
        for j in range(13):ax.text(j,i,f"{heat.iloc[i,j]*100:.1f}",ha="center",va="center",fontsize=7,color="#14221c")
    ax.set(xticks=np.arange(13),xticklabels=[calendar.month_abbr[m] for m in range(1,13)]+["Year"],yticks=np.arange(len(years)),yticklabels=years)
    ax.tick_params(length=0);fig.colorbar(image,ax=ax,shrink=.55,label="Return (%) | color capped at ±12%; cells show actual values")
    save(fig,"04_return_heatmap","Every year, every month","Annual returns are compounded, not the sum of monthly returns. Red = negative; green = positive. Values shown in percent.")

    rates=result["rates"]
    if rates is not None and len(rates.dropna(subset=["stock_return","yield_change_bp"]))>=12:
        fig,axes=plt.subplots(1,2,figsize=(12,5.1));good=rates.dropna(subset=["stock_return","yield_change_bp"])
        axes[0].scatter(good.yield_change_bp,good.stock_return,s=22,alpha=.55,color=GREEN,edgecolors="none")
        r=corr(good.stock_return,good.yield_change_bp)
        axes[0].set(xlabel="Monthly 10-year yield change (bp)",ylabel="Monthly stock return",title=f"Pearson r = {r:+.2f} | n={len(good)}")
        axes[0].axvline(0,lw=.7,color=MUTED);pct(axes[0])
        axes[1].plot(rates.index.to_timestamp("M"),rates.rolling36,color=GREEN,lw=2)
        axes[1].axhline(0,color=MUTED,lw=.8)
        axes[1].set(ylim=(-1,1),title="36-month rolling correlation",ylabel="Correlation");axes[1].grid(axis="y",alpha=.4)
        save(fig,"05_yield_correlation","Optional cross-check | stock returns vs Treasury yield changes","Contemporaneous association only. Uses yield changes, not bond returns, and does not identify a causal or predictive effect.")
    pdf.close()
    return names


def pct_text(x):
    return "n/a" if x is None or not np.isfinite(x) else f"{x:+.2%}"


def safe_json(value):
    if isinstance(value,dict): return {str(k):safe_json(v) for k,v in value.items()}
    if isinstance(value,(list,tuple)): return [safe_json(v) for v in value]
    if isinstance(value,(pd.Timestamp,Path)): return str(value)
    if isinstance(value,(float,np.floating)): return float(value) if np.isfinite(value) else None
    if isinstance(value,np.integer): return int(value)
    return value


def export_excel(result, config, out, methodology):
    """Write a self-contained Excel workbook of chart-ready numerical tables (no embedded charts)."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.properties import CalcProperties

    monthly = result["monthly"]
    yearly = result["yearly"]
    years = result["years"]
    # The numbers below are values (never dependent on Excel's calculation cache).
    # Each metric is calculated from exactly one observation per completed year.
    season = pd.DataFrame([{
        "Month": calendar.month_abbr[m], "Month_Number": m,
        "Sample_Years": len(monthly[m].dropna()),
        "Average_Return": float(monthly[m].mean()),
        "Median_Return": float(monthly[m].median()),
        "Positive_Frequency": float((monthly[m] > 0).mean()),
        "Worst_Return": float(monthly[m].min()),
        "Best_Return": float(monthly[m].max()),
    } for m in range(1,13)])
    annual_p = result["annual_paths"]
    novdec_p = result["seasonal_paths"]
    recent_rows = [i for i, y in enumerate(years) if y >= config["last_year"]-9]
    full_path = pd.DataFrame({
        "Calendar_Day": np.arange(366),
        "Mean_Return": annual_p.mean(axis=0),
        "Median_Return": np.median(annual_p,axis=0),
        "Lower_Quartile": np.quantile(annual_p,.25,axis=0),
        "Upper_Quartile": np.quantile(annual_p,.75,axis=0),
        "Last10_Mean": annual_p[recent_rows].mean(axis=0),
    })
    novdec_path = pd.DataFrame({
        "Days_After_Oct_End": np.arange(62),
        "Mean_Return": novdec_p.mean(axis=0),
        "Median_Return": np.median(novdec_p,axis=0),
        "Lower_Quartile": np.quantile(novdec_p,.25,axis=0),
        "Upper_Quartile": np.quantile(novdec_p,.75,axis=0),
    })
    # One self-contained numeric workbook, without embedded images/charts or external links.
    filepath = out / "results.xlsx"
    with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
        result["daily"].to_excel(writer, sheet_name="Raw_Daily")
        monthly.rename(columns={m:calendar.month_abbr[m] for m in range(1,13)}).to_excel(writer,sheet_name="Monthly_Returns")
        yearly.to_excel(writer,sheet_name="NovDec_By_Year")
        season.to_excel(writer,sheet_name="Monthly_Stats",index=False)
        full_path.to_excel(writer,sheet_name="Full_Year_Path",index=False)
        novdec_path.to_excel(writer,sheet_name="NovDec_Path",index=False)
        result["coverage"].to_excel(writer,sheet_name="Coverage",index=False)
        summary_rows=[]
        for key in ["November","December","NovDec","Last10_NovDec","NovDec_net"]:
            s=result["stats"].get(key,{"n":0})
            if s.get("n",0):
                summary_rows.append({"Period":key,"N":s["n"],"Mean_Return":s["mean"],
                                     "Median_Return":s["median"],"Positive_Frequency":s["win_rate"],
                                     "Worst_Return":s["min"],"Best_Return":s["max"]})
        pd.DataFrame(summary_rows).to_excel(writer,sheet_name="Summary",index=False)
        pd.DataFrame({"Methodology":methodology.splitlines()}).to_excel(writer,sheet_name="Methodology",index=False)
        instructions=[
            ("Research instrument",config["tickers"]["Stock"]),
            ("Market data field",FIELD),
            ("Source",config["source"]),
            ("Return basis",config["basis"]),
            ("Historical sample",f"{min(years)}–{max(years)}, {len(years)} valid complete calendar years"),
            ("Nov-Dec holding period","October last trading close through December last trading close"),
            ("Raw data","Raw_Daily: source daily closing price levels; no missing-data fill"),
            ("Monthly returns","Monthly_Returns: prior month's closing price to current month's closing price"),
            ("Historical observations","NovDec_By_Year: November, December, combined, yearly drawdown"),
            ("Charts / images","Separate PNG and SVG figures are exported outside Excel."),
            ("Replot in Excel","Monthly_Stats, Full_Year_Path and NovDec_Path have chart-ready numeric columns"),
            ("Check calculations","Formula_Audit: Excel formulas for all 12 monthly averages and positive frequencies"),
            ("External links","None: workbook calculations use only local worksheets"),
            ("Statistical caveat","Seasonality is historical and is not a forecast or guaranteed trade signal"),
        ]
        pd.DataFrame(instructions,columns=["Topic","Description"]).to_excel(writer,sheet_name="ReadMe",index=False)
        wb=writer.book
        wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
        audit=wb.create_sheet("Formula_Audit")
        audit.append(["Month","Average_Return_Excel","Average_Return_Python","Difference",
                      "Win_Rate_Excel","Win_Rate_Python","Difference"])
        for i in range(2,14):
            column = get_column_letter(i)  # Monthly_Returns: Year in A, Jan in B ... Dec in M
            audit.append([
                calendar.month_abbr[i-1],
                f"=AVERAGE(Monthly_Returns!{column}2:{column}{1+len(years)})",
                float(season.loc[i-2,"Average_Return"]),
                f"=B{i}-C{i}",
                f'=COUNTIF(Monthly_Returns!{column}2:{column}{1+len(years)},">0")/COUNT(Monthly_Returns!{column}2:{column}{1+len(years)})',
                float(season.loc[i-2,"Positive_Frequency"]),
                f"=E{i}-F{i}",
            ])
        audit.freeze_panes="B2"
        percentage_sheets={
            "Monthly_Returns":list(range(2,14)),
            "NovDec_By_Year":list(range(2,yearly.shape[1]+2)),
            "Monthly_Stats":[4,5,6,7,8],
            "Full_Year_Path":[2,3,4,5,6],
            "NovDec_Path":[2,3,4,5],
            "Summary":[3,4,5,6,7],
            "Formula_Audit":[2,3,4,5,6,7],
        }
        for ws in wb:
            ws.sheet_view.showGridLines=False
            ws.freeze_panes="B2"
            ws.auto_filter.ref=ws.dimensions
            for cell in ws[1]:
                cell.fill=PatternFill("solid",fgColor="173B35")
                cell.font=Font(color="FFFFFF",bold=True)
                cell.alignment=Alignment(wrap_text=True,vertical="center")
            ws.row_dimensions[1].height=30
            for c in ws.columns:
                colletter=c[0].column_letter
                ws.column_dimensions[colletter].width=20 if c[0].column>1 else 24
            for col in percentage_sheets.get(ws.title,[]):
                for r in ws.iter_rows(min_row=2,min_col=col,max_col=col):
                    for cell in r:
                        cell.number_format='0.00%;[Red](0.00%);–'
            if ws.title=="ReadMe":
                ws.column_dimensions["A"].width=29
                ws.column_dimensions["B"].width=110
            if ws.title=="Methodology":ws.column_dimensions["A"].width=125
            if ws.title=="Raw_Daily":
                for r in ws.iter_rows(min_row=2,min_col=1,max_col=1):r[0].number_format="yyyy-mm-dd"
        # Keep the workbook data-first. Separate PNG / SVG charts are produced by figures().
        writer.book.active=0
    # Verify workbook archive integrity and ensure it contains no native charts.
    from zipfile import ZipFile
    with ZipFile(filepath) as archive:
        if archive.testzip() is not None: raise RuntimeError("The exported Excel ZIP is damaged")
        if any(n.startswith("xl/charts/") for n in archive.namelist()):
            raise RuntimeError("Unexpected native Excel chart content")


def export_report(result, config, names, out):
    s=result["stats"];years=result["years"]
    metadata={"settings":config,"accepted_years":years,"statistics":s,
              "raw_data_sha256":hashlib.sha256((out/"raw_daily.csv").read_bytes()).hexdigest()}
    (out/"summary.json").write_text(json.dumps(safe_json(metadata),indent=2,ensure_ascii=False,allow_nan=False),encoding="utf-8")
    method=("Return basis: "+config["basis"]+". Input levels are never automatically converted to total return.\n"
            "Complete calendar years only; daily history quality checked, no missing-month fill or ETF/index splice.\n"
            "Each year equal weight. November-December = Dec last close / Oct last close - 1.\n"
            "Calendar-day curves forward-fill between genuine observations only after full-year quality checks.\n"
            "Leap day removed from the display grid; its economic return remains in subsequent cumulative levels.\n"
            "Shaded path bands are cross-year quartiles, not confidence intervals.\n"
            "Bootstrap resamples whole yearly observations; assumes exchangeable years and is not a future prediction interval.\n"
            "IWM price returns exclude cash distributions; no dividends, trading costs (unless set), tax or cash yield.\n"
            "Optional round-trip cost is split equally at entry/exit: (1-c/2)^2*(1+return)-1. Charts are gross.\n"
            "Optional US10y is in percent units; monthly yield differences times 100 are basis points.\n"
            "Report all missing history and actual N. No multiple-testing correction or causal identification.\n")
    (out/"methodology.txt").write_text(method,encoding="utf-8")
    # Full Excel deliverable is mandatory; Excel contains numeric tables, not charts.
    export_excel(result, config, out, method)
    rows=[]
    for key,title in [("NovDec","Nov–Dec"),("November","November"),("December","December"),
                      ("Last10_NovDec","Last 10 calendar years: Nov–Dec"),("Excess_NovDec","Nov–Dec stock minus benchmark")]:
        a=s.get(key,{"n":0})
        if not a["n"]:continue
        rows.append(f"<tr><td>{title}</td><td>{a['n']}</td><td>{pct_text(a['mean'])}</td><td>{pct_text(a['median'])}</td><td>{a['win_rate']:.0%}</td><td>{pct_text(a['min'])}</td><td>{pct_text(a['max'])}</td></tr>")
    images=[]
    for name,title in names:
        data=base64.b64encode((out/(name+".svg")).read_bytes()).decode()
        images.append(f'<section><h2>{html.escape(title)}</h2><img src="data:image/svg+xml;base64,{data}" alt="{html.escape(title)}"><p class="small"><a href="{name}.png">PNG</a> · <a href="{name}.svg">SVG</a></p></section>')
    top=s["NovDec"];ci=top["mean_bootstrap_95ci"]
    demo='<div class="warning">SYNTHETIC DEMO — NOT BLOOMBERG DATA. Do not use these numbers for investment decisions.</div>' if config["demo"] else ''
    missing=''.join('<li>'+html.escape(w)+'</li>' for w in config["warnings"])
    report=f'''<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Russell 2000 | Seasonality</title><style>
    *{{box-sizing:border-box}}body{{margin:0;background:{PAPER};color:{INK};font:15px/1.8 Arial,sans-serif}}main{{max-width:1160px;margin:auto;padding:40px}}h1{{font:600 42px/1.3 Georgia,serif;margin:14px 0}}h2{{font-size:22px}}.eyebrow{{color:{GREEN};font-size:11px;letter-spacing:.16em;font-weight:bold}}.small{{color:#5c6862;font-size:12px}}section{{padding:24px 0;border-top:1px solid #d6dcd6}}img{{width:100%;height:auto}}.scroll{{overflow:auto}}table{{width:100%;border-collapse:collapse;min-width:660px}}th,td{{text-align:left;padding:12px 10px;border-bottom:1px solid #d6dcd6;font-size:13px}}th{{color:#5c6862}}a{{color:{GREEN}}}.note{{border-left:2px solid {GREEN};padding-left:17px}}.warning{{background:#f4dfd8;padding:18px;color:#853b30}}button{{background:none;border:1px solid #cbd3cb;padding:8px 13px;cursor:pointer;float:right}}pre{{white-space:pre-wrap;overflow-wrap:anywhere;font:12px/1.8 Arial,sans-serif}}@media(max-width:700px){{main{{padding:20px}}h1{{font-size:30px}}}}@media print{{@page{{size:A4;margin:13mm}}main{{padding:0}}body{{background:white;font-size:10pt}}button{{display:none}}img{{break-inside:avoid}}section{{padding:16px 0}}table{{min-width:0}}th,td{{font-size:8pt;padding:7px}}.scroll{{overflow:visible}}}}
    </style></head><body><main><button onclick="window.print()">Print / PDF</button><div class="eyebrow">SEASONALITY RESEARCH</div><h1>IWM ETF<br><span style="color:{GREEN}">{len(years)} complete years</span></h1>{demo}
    <p>{html.escape(config['tickers']['Stock'])} · {html.escape(config['basis'])}<br>Requested: {config['first_year']}–{config['last_year']} ({config['years']} years). Accepted: <b>{len(years)} years</b>.</p>
    <p class="small">Research date: {config['as_of']} · Source: {html.escape(config['source'])} · Field: {html.escape(FIELD)}</p>
    <section><h2>November–December at a glance</h2><p>Mean <b>{pct_text(top['mean'])}</b> · Median <b>{pct_text(top['median'])}</b> · Positive <b>{top['win_rate']:.0%}</b> · n={top['n']}</p>
    <div class="scroll"><table><thead><tr><th>Window</th><th>N</th><th>Mean</th><th>Median</th><th>Positive</th><th>Worst</th><th>Best</th></tr></thead><tbody>{''.join(rows)}</tbody></table></div>
    <p class="note">95% bootstrap interval for the Nov–Dec MEAN: {pct_text(ci[0])} to {pct_text(ci[1])}. This is not a forecast range for next year. Annual observations are not guaranteed to be identically distributed.</p>
    <p>Worst observed within-window drawdown: <b>{pct_text(result['yearly'].NovDec_max_drawdown.min())}</b>. A favorable average does not rule out a severe interim loss.</p>
    <p class="small">Returns in tables and charts are gross. Optional round-trip cost: {ROUND_TRIP_COST_BPS:g}bp; after-cost returns are available in nov_dec_by_year.csv. Missing years are never synthesized.</p></section>
    {''.join(images)}<section><h2>Data quality and reproducibility</h2><ul>{missing or '<li>All requested stock years passed the configured daily-history checks.</li>'}</ul><p class="small">Accepted years: {', '.join(map(str,years))}. Benchmark and rate samples may be shorter; each statistic shows its own N.</p>
    <p><a href="raw_daily.csv">Raw daily data</a> · <a href="coverage.csv">Coverage</a> · <a href="monthly_returns.csv">Monthly returns</a> · <a href="nov_dec_by_year.csv">Annual Nov–Dec returns</a> · <a href="summary.json">Statistics / settings</a> · <a href="charts.pdf">Charts PDF</a></p><pre>{html.escape(method)}</pre></section></main></body></html>'''
    (out/"report.html").write_text(report,encoding="utf-8")
    # Rewrite after optional-export warnings so JSON and HTML describe the same run.
    (out/"summary.json").write_text(json.dumps(safe_json(metadata),indent=2,ensure_ascii=False,allow_nan=False),encoding="utf-8")


def demo_data(config):
    dates=pd.bdate_range(config["start"],config["end"])
    rng=np.random.default_rng(RANDOM_SEED)
    shocks=rng.normal(.0002,.012,len(dates))
    return pd.DataFrame({"Stock":100*np.exp(np.cumsum(shocks)),
                         "Benchmark":100*np.exp(np.cumsum(.75*shocks+rng.normal(.00005,.004,len(dates)))),
                         "US10y":np.clip(4+np.cumsum(rng.normal(0,.025,len(dates))),.1,15)},index=dates)


def self_test():
    """Deterministic financial edge cases; never calls Bloomberg."""
    from types import SimpleNamespace
    args=SimpleNamespace(as_of="2022-10-25",years=2,allow_short_history=False,demo=True)
    cfg=settings(args)
    assert cfg["first_year"]==2020 and cfg["last_year"]==2021
    dates=pd.bdate_range("2019-12-01","2021-12-31")
    offsets=(dates.year-2019)*12+dates.month-12
    sample=pd.DataFrame({"Stock":100*1.01**offsets,"Benchmark":100*1.005**offsets},index=dates)
    with tempfile.TemporaryDirectory(prefix="iwm_test_") as folder:
        res=analyze(sample,cfg,Path(folder))
        assert res["years"]==[2020,2021]
        assert res["annual_paths"].shape==(2,366) and res["seasonal_paths"].shape==(2,62)
        assert np.allclose(res["monthly"],.01)
        assert np.allclose(res["yearly"].NovDec,1.01**2-1)
        assert np.allclose(res["yearly"].FullYear,1.01**12-1)
        assert np.allclose(res["seasonal_paths"][:,-1],res["yearly"].NovDec)
        assert (res["yearly"].NovDec_max_drawdown==0).all()
        broken=sample.drop(sample.loc["2020-06"].index)
        good,coverage,_=validate_years(broken.Stock,cfg)
        assert good==[2021] and not coverage.iloc[0].valid
        try: analyze(broken,cfg,Path(folder))
        except ValueError: pass
        else: raise AssertionError("Missing month must fail strict history validation")
        short=sample.loc["2020-01":]
        assert validate_years(short.Stock,cfg)[0]==[2021]  # missing baseline December
    wide=pd.DataFrame({("IWM US Equity",FIELD):[100.,101.]},index=pd.to_datetime(["2020-01-02","2020-01-03"]))
    assert normalize_bdh(wide,"IWM US Equity").tolist()==[100.,101.]
    long=pd.DataFrame({"date":["2020-01-02"],"ticker":["IWM US Equity"],"field":[FIELD],"value":[100.]})
    assert normalize_bdh(long,"IWM US Equity").iloc[0]==100
    assert abs(corr(pd.Series([1,2,3]),pd.Series([3,2,1]))+1)<1e-12
    assert stats([.1,-.1])["mean"]==0
    try:clean_daily(pd.DataFrame({"Stock":[100,101]},index=["2020-01-02"]*2),{"Stock":"IWM US Equity"})
    except ValueError:pass
    else:raise AssertionError("Duplicate dates must fail")
    print("SELF-TEST PASSED: compounding, leap years, October weekend baseline, missing months, missing baseline, duplicate dates, Bloomberg layouts, correlations.")


def main():
    parser=argparse.ArgumentParser(description=__doc__,formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--years",type=int,default=YEARS)
    parser.add_argument("--as-of",help="Research date; current year is excluded")
    parser.add_argument("--backend",choices=["auto","xbbg","pdblp","blpapi"],default=BACKEND)
    parser.add_argument("--csv",type=Path,help="Existing wide daily-level data; no Bloomberg required")
    parser.add_argument("--out",type=Path)
    parser.add_argument("--allow-short-history",action="store_true")
    parser.add_argument("--demo",action="store_true",help="Synthetic preview, prominently watermarked")
    parser.add_argument("--self-test",action="store_true",help="Offline calculation tests")
    args=parser.parse_args()
    logging.basicConfig(level=logging.INFO,format="%(levelname)s | %(message)s")
    if args.self_test:self_test();return
    if args.demo and args.csv:parser.error("Choose either --demo or --csv")
    # Fail fast if a required Excel/plotting dependency is unavailable.
    import openpyxl
    import matplotlib
    cfg=settings(args)
    out=args.out or (Path(OUTPUT_DIR) if OUTPUT_DIR else Path(__file__).resolve().parent/("demo_output" if args.demo else "output"))
    out=out.expanduser().resolve();out.mkdir(parents=True,exist_ok=True)
    # Never leave an old report looking like the outcome of a failed new run.
    if (out/"summary.json").exists() or (out/"report.html").exists():
        out=out/ ("run_"+pd.Timestamp.now().strftime("%Y%m%d_%H%M%S_%f"));out.mkdir(parents=True)
    cfg["output_directory"]=str(out)
    LOG.info("Sample: %s–%s, %s requested full years",cfg["first_year"],cfg["last_year"],cfg["years"])
    if args.demo:
        cfg["source"]="SYNTHETIC DEMO — NOT BLOOMBERG";raw=demo_data(cfg)
    elif args.csv:
        cfg["source"]="User CSV: "+args.csv.name
        cfg["warnings"].append("CSV index identifiers / return basis are user-provided settings, not independently verified.")
        raw=pd.read_csv(args.csv,index_col=0)
    else:raw=pull_bloomberg(cfg,args.backend)
    daily=clean_daily(raw,cfg["tickers"])
    for key in [k for k in cfg["tickers"] if k not in daily]:
        cfg["warnings"].append(key+" absent from input; its comparison is omitted.")
    result=analyze(daily,cfg,out)
    names=figures(result,cfg,out)
    export_report(result,cfg,names,out)
    a=result["stats"]["NovDec"]
    print("\n"+("SYNTHETIC DEMO | " if args.demo else "")+f"Nov–Dec: N={a['n']}, mean={a['mean']:+.2%}, median={a['median']:+.2%}, positive={a['win_rate']:.0%}")
    print("Return basis: "+cfg["basis"])
    print("Open: "+str(out/"report.html"))
    print("Charts: "+str(out/"charts.pdf"))
    if cfg["warnings"]:print("Warnings are listed in report.html and summary.json.")


if __name__ == "__main__":
    try:main()
    except Exception as error:
        LOG.exception("Seasonality run failed: %s",error)
        LOG.error("No automatic data substitution. Fix the source/settings and rerun. Use --self-test for offline checks.")
        sys.exit(1)
